"""Synthetic end-to-end checks of PRIMARY assembly, mapping, validation and export."""

from __future__ import annotations

import copy
import io
import json
import sqlite3
from pathlib import Path

import openpyxl
import pytest

from bctc_ai.primary_statements import catalog_v1
from bctc_ai.primary_statements.catalog_v1 import PrimaryCatalogError, load_primary_catalog_v1
from bctc_ai.primary_statements.excel_v1 import (
    PrimaryWorkbookError,
    build_workbook_v1,
    excel_number_v1,
    template_rows_v1,
    verify_workbook_v1,
)
from bctc_ai.primary_statements.pipeline_v1 import run_document_v1
from bctc_ai.primary_statements.result_v1 import build_document_result_v1
from bctc_ai.primary_statements.source_v1 import (
    PrimarySourceError,
    SelectedDocument,
    SelectedPage,
    load_corpus_selection_v1,
)
from bctc_ai.primary_statements.store_v1 import PrimaryStoreError, reload_result_store_v1, write_result_store_v1
from bctc_ai.primary_statements.validation_v1 import load_equation_config_v1

ROOT = Path(__file__).resolve().parents[2]


@pytest.fixture(scope="module")
def catalog():
    return load_primary_catalog_v1(ROOT)


@pytest.fixture(scope="module")
def equations():
    return load_equation_config_v1(ROOT)


def row(label, values, path=None, kind="ITEM"):
    return {"label_exact": label, "row_kind": kind, "hierarchy_path_exact": path or [label], "values_exact": values}


def table(rows, headers, *, title=None, unit=None, text_columns=()):
    columns = [{"header_path_exact": [h], "value_kind": "TEXT"} for h in text_columns]
    columns += [{"header_path_exact": h if isinstance(h, list) else [h], "value_kind": "MONEY"} for h in headers]
    return {"title_exact": title, "unit_exact": unit, "continuation": "NONE", "columns": columns, "rows": rows}


def section(tables, *, statement="BALANCE_SHEET", kind="PRIMARY_STATEMENT", title=None, narratives=()):
    return {"content_kind": kind, "statement_type": statement, "title_exact": title,
            "narratives_exact": list(narratives), "tables": tables}


def run_doc(catalog, equations, pages):
    selected = []
    census, loaded = {}, {}
    for number, sections in enumerate(pages, 1):
        vid = "gfpstorev1:json:" + f"{number:064x}"
        selected.append(SelectedPage(number, vid, f"page-{number}", "PRIMARY_FINANCIAL_STATEMENT", "0" * 64))
        census[vid] = {"candidate": True, "primary_metadata": True, "max_hits": {}, "tables": 1,
                       "census_reason": "PRIMARY_METADATA"}
        loaded[vid] = {"page_json": {"status": "PRIMARY_FINANCIAL_STATEMENT", "sections": sections}}
    document = SelectedDocument(1, "synthetic.pdf", "a" * 64, 10, len(pages), "gjfpdocv1:" + "b" * 64,
                                "gfdmv1:manifest:" + "c" * 64, "gjfcdmsv1:selection:" + "d" * 64, tuple(selected))
    return run_document_v1(document, census, {}, loaded, catalog, ROOT / "config/mapping/lctt-v2.yaml", equations)


def result_of(run, catalog):
    return build_document_result_v1(run, catalog, corpus_identity={"corpus_name": "synthetic"}, pins={"test": True})


BS_TITLE = "BÁO CÁO TÌNH HÌNH TÀI CHÍNH HỢP NHẤT\nTại ngày 30 tháng 6 năm 2025"
BS_HEADERS = [["30/06/2025", "Triệu VND"], ["31/12/2024", "Triệu VND"]]


def balance_sheet_rows(cash=("100", "90")):
    return [
        row("A TÀI SẢN", [None, None], ["A TÀI SẢN"], "GROUP"),
        row("I Tiền mặt, vàng bạc, đá quý", list(cash), ["A TÀI SẢN", "I Tiền mặt"]),
        row("III Tiền gửi và cho vay các TCTD khác", ["50", "40"], ["A TÀI SẢN", "III Tiền gửi"]),
        row("1 Tiền gửi tại các TCTD khác", ["55", "44"], ["A TÀI SẢN", "III Tiền gửi", "1"]),
        row("2 Dự phòng rủi ro", ["(5)", "(4)"], ["A TÀI SẢN", "III Tiền gửi", "2"]),
        row("VI Cho vay khách hàng", ["900", "800"], ["A TÀI SẢN", "VI Cho vay khách hàng"]),
        row("1 Cho vay khách hàng", ["910", "805"], ["A TÀI SẢN", "VI Cho vay khách hàng", "Cho vay khách hàng"]),
        row("2 Dự phòng rủi ro cho vay khách hàng", ["(10)", "(5)"], ["A TÀI SẢN", "VI Cho vay khách hàng", "Dự phòng"]),
        row("TỔNG TÀI SẢN", ["1.050", "930"], ["TỔNG TÀI SẢN"], "TOTAL"),
    ]


def bs_page(rows, *, title=BS_TITLE, headers=BS_HEADERS, kind="PRIMARY_STATEMENT"):
    return [section([table(rows, headers, title=title)], title=title, kind=kind)]


def cells_by(result, rnid, period_key=None):
    return [c for c in result["cells"] if c["report_norm_id"] == rnid and (period_key is None or c["period_key"] == period_key)]


def test_balance_sheet_structure_mapping_and_acceptance(catalog, equations):
    run = run_doc(catalog, equations, [bs_page(balance_sheet_rows())])
    result = result_of(run, catalog)
    [instance] = result["instances"]
    assert instance["scope"] == "CONSOLIDATED" and instance["status"] == "ASSEMBLED"
    ids = {row["label_exact"]: row["report_norm_id"] for row in instance["source_rows"]}
    assert ids["VI Cho vay khách hàng"] == 4315 and ids["1 Cho vay khách hàng"] == 4348
    assert ids["2 Dự phòng rủi ro"] == 4345 and ids["1 Tiền gửi tại các TCTD khác"] == 4344
    assert [lane["period"]["key"] for lane in instance["lanes"]] == ["INSTANT:2025-06-30", "INSTANT:2024-12-31"]
    assert all(lane["unit"] == "MILLION_VND" for lane in instance["lanes"])
    statuses = {check["equation_id"]: check["status"] for check in instance["validations"] if check["lane"] == 0}
    assert statuses["CDKT:4315"] == "PASS" and statuses["CDKT:4312"] == "PASS"
    loan = cells_by(result, 4348, "INSTANT:2025-06-30")[0]
    assert loan["acceptance"] == "ACCEPTED" and loan["value"] == "910" and loan["locator"]["row_ordinal"] == 7
    items = {i["report_norm_id"]: i["status"] for i in result["item_status"]}
    assert items[4315] == "MAPPED" and items[4311] == "NOT_OBSERVED"


def test_note_table_with_statement_labels_is_not_primary(catalog, equations):
    run = run_doc(catalog, equations, [bs_page(balance_sheet_rows(), title="Rủi ro tiền tệ", kind="FINANCIAL_NOTE")])
    assert not run.instances
    assert run.other_tables[0].classification["decision"] == "NOT_PRIMARY"


def test_duplicate_target_with_different_values_is_unresolved_not_first_wins(catalog, equations):
    rows = balance_sheet_rows()
    rows.insert(2, row("Tiền mặt, vàng bạc, đá quý", ["101", "90"], ["A TÀI SẢN", "Tiền mặt x"]))
    result = result_of(run_doc(catalog, equations, [bs_page(rows)]), catalog)
    dispositions = [r["disposition"] for r in result["instances"][0]["source_rows"] if r["label_key"] == "tien mat vang bac da quy"]
    assert dispositions == ["UNRESOLVED", "UNRESOLVED"]
    assert not cells_by(result, 4310)


def test_identical_carry_forward_row_is_control_and_not_a_restart(catalog, equations):
    kq_title = "BÁO CÁO KẾT QUẢ HOẠT ĐỘNG HỢP NHẤT\nCho kỳ sáu tháng kết thúc ngày 30 tháng 6 năm 2025"
    headers = [["Kỳ sáu tháng kết thúc ngày 30/6/2025", "Triệu VND"], ["Kỳ sáu tháng kết thúc ngày 30/6/2024", "Triệu VND"]]
    page1 = [row("Thu nhập lãi và các khoản thu nhập tương tự", ["10", "9"]),
             row("Chi phí lãi và các chi phí tương tự", ["(4)", "(3)"]),
             row("Thu nhập lãi thuần", ["6", "6"]),
             row("Tổng lợi nhuận trước thuế (mang sang trang sau)", ["6", "6"])]
    page2 = [row("Tổng lợi nhuận trước thuế (mang sang từ trang trước)", ["6", "6"]),
             row("Chi phí thuế TNDN hiện hành", ["(1)", "(1)"]),
             row("Lợi nhuận sau thuế", ["5", "5"])]
    pages = [[section([table(page1, headers, title=kq_title)], statement="INCOME_STATEMENT", title=kq_title)],
             [section([table(page2, headers, title=kq_title)], statement="INCOME_STATEMENT", title=kq_title)]]
    result = result_of(run_doc(catalog, equations, pages), catalog)
    assert len(result["instances"]) == 1
    pbt = [r["disposition"] for r in result["instances"][0]["source_rows"] if r["report_norm_id"] == 4377 or r["label_key"] == "tong loi nhuan truoc thue"]
    assert sorted(pbt) == ["DUPLICATE_IDENTICAL_CONTROL", "MAPPED"]
    assert len(cells_by(result, 4377)) == 2


def test_vnd_and_million_presentations_retire_control_with_cell_veto(catalog, equations):
    vnd = [row("Tiền mặt, vàng bạc, đá quý", ["100.400.000", "90.000.000"]),
           row("Tiền gửi tại NHNN", ["5.000.000", "4.000.000"]),
           row("Cho vay khách hàng", ["900.000.000", "800.000.000"])]
    mil = [row("Tiền mặt, vàng bạc, đá quý", ["100", "90"]),
           row("Tiền gửi tại NHNN", ["5", "4"]),
           row("Cho vay khách hàng", ["903", "800"])]
    vnd_headers = [["30/06/2025", "VND"], ["31/12/2024", "VND"]]
    pages = [bs_page(vnd, headers=vnd_headers), bs_page(mil)]
    run = run_doc(catalog, equations, pages)
    receipt = run.presentation_receipts[0]
    assert receipt["status"] == "CONTROL_RETIRED_WITH_CELL_VETOES" and receipt["mismatch_count"] == 1
    result = result_of(run, catalog)
    held = cells_by(result, 4315, "INSTANT:2025-06-30")[0]
    assert held["acceptance"] == "HOLD" and any("CROSS_PRESENTATION" in r for r in held["hold_reasons"])
    assert cells_by(result, 4310, "INSTANT:2025-06-30")[0]["acceptance"] == "ACCEPTED"
    assert all(c["unit"] == "VND" for c in result["cells"])


def test_blank_is_never_zero_and_contradicted_dash_is_not_a_zero(catalog, equations):
    rows = balance_sheet_rows()
    rows[1] = row("I Tiền mặt, vàng bạc, đá quý", [None, "90"], ["A TÀI SẢN", "I Tiền mặt"])
    rows[5] = row("VI Cho vay khách hàng", ["-", "800"], ["A TÀI SẢN", "VI Cho vay khách hàng"])
    result = result_of(run_doc(catalog, equations, [bs_page(rows)]), catalog)
    blank = cells_by(result, 4310, "INSTANT:2025-06-30")[0]
    assert blank["observation_state"] == "BLANK" and blank["value"] is None and blank["acceptance"] == "ACCEPTED_BLANK"
    dash = cells_by(result, 4315, "INSTANT:2025-06-30")[0]
    assert dash["observation_state"] == "DASH" and dash["value"] is None and dash["acceptance"] == "HOLD"


def test_unit_and_period_must_be_printed(catalog, equations):
    rows = balance_sheet_rows()
    result = result_of(run_doc(catalog, equations, [bs_page(rows, title="BÁO CÁO TÌNH HÌNH TÀI CHÍNH HỢP NHẤT",
                                                           headers=["Số cuối kỳ", "Số đầu năm"])]), catalog)
    assert {c["acceptance"] for c in result["cells"]} == {"HOLD"}
    reasons = {r.split(":")[0] for c in result["cells"] for r in c["hold_reasons"]}
    assert {"PERIOD_UNRESOLVED", "UNIT_UNRESOLVED"} <= reasons


def test_scope_from_printed_form_code_and_unresolved_scope_holds(catalog, equations):
    title = "BÁO CÁO TÌNH HÌNH TÀI CHÍNH\nTại ngày 30 tháng 6 năm 2025"
    page = [section([table(balance_sheet_rows(), BS_HEADERS, title=title)], title=title, narratives=["Mẫu B02a/TCTD"])]
    assert result_of(run_doc(catalog, equations, [page]), catalog)["instances"][0]["scope"] == "SEPARATE"
    page = [section([table(balance_sheet_rows(), BS_HEADERS, title=title)], title=title)]
    result = result_of(run_doc(catalog, equations, [page]), catalog)
    assert result["instances"][0]["scope"] is None
    assert all("SCOPE_UNRESOLVED" in c["hold_reasons"] for c in result["cells"])


def test_cash_flow_branch_from_ordered_anchors_and_mixed_branch_holds(catalog, equations):
    title = "BÁO CÁO LƯU CHUYỂN TIỀN TỆ HỢP NHẤT\nCho kỳ sáu tháng kết thúc ngày 30 tháng 6 năm 2025"
    headers = [["Kỳ sáu tháng kết thúc ngày 30/6/2025", "Triệu VND"]]
    direct = [row("Thu nhập lãi và các khoản thu nhập tương tự nhận được", ["10"]),
              row("Chi phí lãi và các chi phí tương tự đã trả", ["(4)"]),
              row("Lưu chuyển tiền thuần từ hoạt động kinh doanh", ["6"])]
    run = run_doc(catalog, equations, [[section([table(direct, headers, title=title)], statement="CASH_FLOW", title=title)]])
    result = result_of(run, catalog)
    assert result["instances"][0]["cash_flow_branch"] == "DIRECT"
    assert cells_by(result, 4110) and not cells_by(result, 4163)
    statuses = {i["report_norm_id"]: i["status"] for i in result["item_status"]}
    assert statuses[4163] == "NOT_APPLICABLE"
    mixed = direct + [row("Lợi nhuận trước thuế", ["7"]), row("Điều chỉnh cho các khoản", [None])]
    run = run_doc(catalog, equations, [[section([table(mixed, headers, title=title)], statement="CASH_FLOW", title=title)]])
    result = result_of(run, catalog)
    assert result["instances"][0]["cash_flow_branch"] == "BOTH"
    assert all(c["acceptance"] == "HOLD" for c in result["cells"])


def test_quarter_and_ytd_lanes_keep_distinct_write_keys(catalog, equations):
    title = "BÁO CÁO KẾT QUẢ HOẠT ĐỘNG HỢP NHẤT\nQuý 2 năm 2025"
    headers = [["Quý 2", "Năm nay"], ["Quý 2", "Năm trước"],
               ["Lũy kế từ đầu năm đến cuối quý này", "Năm nay"], ["Lũy kế từ đầu năm đến cuối quý này", "Năm trước"]]
    rows = [row("Thu nhập lãi và các khoản thu nhập tương tự", ["10", "9", "20", "18"])]
    page = [section([table(rows, headers, title=title, unit="Đơn vị: triệu đồng")], statement="INCOME_STATEMENT", title=title)]
    result = result_of(run_doc(catalog, equations, [page]), catalog)
    keys = [c["period_key"] for c in result["cells"]]
    assert keys == ["DURATION:QUARTER:2:2025", "DURATION:QUARTER:2:PRIOR_YEAR_OF(2025)",
                    "DURATION:YTD:Q2:2025", "DURATION:YTD:Q2:PRIOR_YEAR_OF(2025)"]
    assert {c["acceptance"] for c in result["cells"]} == {"ACCEPTED"}


def test_code_column_is_not_ancestry_and_fragments_merge(catalog, equations):
    rows = [row("TÀI SẢN", [None, None, None], ["TÀI SẢN"], "GROUP"),
            row("X", [None, "100", "90"], ["TÀI SẢN", "X"]),
            row("Tài sản cố định", ["120", "100", "90"], ["TÀI SẢN", "X", "Tài sản cố định"]),
            row("1 Tài sản cố định hữu hình", ["121", "100", "90"], ["TÀI SẢN", "Tài sản cố định", "1"]),
            row("a Nguyên giá", ["122", "150", "130"], ["TÀI SẢN", "Tài sản cố định", "1", "a"]),
            row("b Giá trị hao mòn lũy kế", ["123", "(50)", "(40)"], ["TÀI SẢN", "Tài sản cố định", "1", "b"])]
    page = [section([table(rows, BS_HEADERS, title=BS_TITLE, text_columns=("Mã số",))], title=BS_TITLE)]
    result = result_of(run_doc(catalog, equations, [page]), catalog)
    source_rows = result["instances"][0]["source_rows"]
    by_label = {r["label_exact"]: r for r in source_rows}
    assert by_label["Tài sản cố định"]["merged_fragments"][0]["label_exact"] == "X"
    assert by_label["a Nguyên giá"]["report_norm_id"] == 4367
    assert by_label["b Giá trị hao mòn lũy kế"]["report_norm_id"] == 4368
    assert by_label["a Nguyên giá"]["code"] == "122"


def test_excel_exact_readback_precision_and_tamper(catalog, equations, tmp_path):
    rows = balance_sheet_rows(cash=("9.007.199.254.740.993", "90"))
    result = result_of(run_doc(catalog, equations, [bs_page(rows)]), catalog)
    assert excel_number_v1("9007199254740993") == ("9007199254740993", "EXACT_TEXT_INTEGER_BEYOND_DOUBLE_PRECISION")
    assert excel_number_v1("1112") == (1112, "EXCEL_INTEGER_EXACT")
    template = template_rows_v1(ROOT, catalog)
    payload, _ = build_workbook_v1(result, template)
    assert verify_workbook_v1(payload, result, template)["status"] == "PASS_EXACT_READBACK"
    assert build_workbook_v1(result, template)[0] == payload  # deterministic bytes
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    workbook["CDKT"]["E3"].value = 1
    tampered = io.BytesIO()
    workbook.save(tampered)
    with pytest.raises(PrimaryWorkbookError):
        verify_workbook_v1(tampered.getvalue(), result, template)
    altered = copy.deepcopy(result)
    altered["cells"][0]["unit"] = "VND"
    with pytest.raises(PrimaryWorkbookError):
        verify_workbook_v1(payload, altered, template)


def test_formula_like_source_text_stays_literal(catalog, equations):
    rows = balance_sheet_rows()
    rows.append(row("=SUM(A1:A9)", ["1", "1"], ["=SUM(A1:A9)"]))
    result = result_of(run_doc(catalog, equations, [bs_page(rows)]), catalog)
    template = template_rows_v1(ROOT, catalog)
    payload, _ = build_workbook_v1(result, template)
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    labels = [cell.value for cell in workbook["SOURCE_ROWS"]["E"]]
    assert "=SUM(A1:A9)" in labels
    assert all(cell.data_type != "f" for sheet in workbook for line in sheet.iter_rows() for cell in line)


def test_store_reload_detects_projection_and_document_tamper(catalog, equations, tmp_path):
    result = result_of(run_doc(catalog, equations, [bs_page(balance_sheet_rows())]), catalog)
    store = tmp_path / "store.sqlite3"
    write_result_store_v1(store, run_manifest={"run": 1}, results=[result])
    manifest, reloaded = reload_result_store_v1(store)
    assert reloaded[0]["result_id"] == result["result_id"]
    with pytest.raises(PrimaryStoreError):
        write_result_store_v1(store, run_manifest={"run": 1}, results=[result])
    connection = sqlite3.connect(store)
    connection.execute("UPDATE mapped_cell SET value_text='911' WHERE report_norm_id=4348 AND lane=0")
    connection.commit()
    connection.close()
    with pytest.raises(PrimaryStoreError, match="mapped_cell"):
        reload_result_store_v1(store)
    other = tmp_path / "other.sqlite3"
    write_result_store_v1(other, run_manifest={"run": 1}, results=[result])
    connection = sqlite3.connect(other)
    raw = connection.execute("SELECT result_json FROM document_result").fetchone()[0]
    forged = raw.replace(b'"value":"910"', b'"value":"911"')
    import hashlib
    connection.execute("UPDATE document_result SET result_json=?, result_sha256=?", (forged, hashlib.sha256(forged).hexdigest()))
    connection.commit()
    connection.close()
    with pytest.raises(PrimaryStoreError, match="result id"):
        reload_result_store_v1(other)


def test_store_rejects_duplicate_accepted_write_keys(catalog, equations, tmp_path):
    result = result_of(run_doc(catalog, equations, [bs_page(balance_sheet_rows())]), catalog)
    duplicate = copy.deepcopy(result)
    cell = copy.deepcopy(next(c for c in duplicate["cells"] if c["acceptance"] == "ACCEPTED"))
    cell["lane"] = 9
    duplicate["cells"].append(cell)
    with pytest.raises(sqlite3.IntegrityError):
        write_result_store_v1(tmp_path / "dup.sqlite3", run_manifest={}, results=[duplicate])


def test_stale_or_tampered_corpus_index_is_rejected(tmp_path):
    index = tmp_path / "index.json"
    index.write_text(json.dumps({"format_version": "X"}))
    with pytest.raises(PrimarySourceError, match="SHA-256"):
        load_corpus_selection_v1(name="t", root=tmp_path, index_relative_path="index.json",
                                 expected_index_sha256="0" * 64, expected_db_sha256="0" * 64)


def test_catalog_rejects_unconstrained_alias_collision(tmp_path, monkeypatch):
    original = (ROOT / catalog_v1.CATALOG_RELATIVE_PATH).read_text(encoding="utf-8")
    forged = original.replace('4311:\n        aliases: ["Tiền gửi tại Ngân hàng Nhà nước", "Tiền gửi tại NHNN"]',
                              '4311:\n        aliases: ["Tiền gửi tại Ngân hàng Nhà nước", "Tiền gửi tại NHNN", "Tiền mặt"]')
    assert forged != original
    path = tmp_path / "catalog.yaml"
    path.write_text(forged, encoding="utf-8")
    monkeypatch.setattr(catalog_v1, "CATALOG_RELATIVE_PATH", str(path))
    with pytest.raises(PrimaryCatalogError, match="unconstrained"):
        load_primary_catalog_v1(ROOT)


def test_catalog_change_changes_pinned_identity(tmp_path, monkeypatch, catalog):
    original = (ROOT / catalog_v1.CATALOG_RELATIVE_PATH).read_text(encoding="utf-8")
    path = tmp_path / "catalog.yaml"
    path.write_text(original + "\n# reviewed comment\n", encoding="utf-8")
    monkeypatch.setattr(catalog_v1, "CATALOG_RELATIVE_PATH", str(path))
    changed = load_primary_catalog_v1(ROOT)
    assert changed.catalog_sha256 != catalog.catalog_sha256
