"""Canonical-proposal TM workbook: numbers only from persisted canonical mappings; every ledger cell on DIMENSIONAL."""

from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from test_f36_notes_core_adapter_v1 import example as example

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes
from bctc_ai.notes.export_v1 import NotesWorkbookError, tm_template_rows_v1
from bctc_ai.notes.f36_canonical_workbook_v1 import (
    F36CanonicalWorkbookError,
    build_f36_canonical_workbook_v1,
    verify_f36_canonical_workbook_v1,
)

ROOT = Path(__file__).resolve().parents[2]
DOCUMENT = {"document_plan_id": "plan", "relative_path": "synthetic.pdf", "source_sha256": "1" * 64}
PINS = {"synthetic": "b" * 64}


@pytest.fixture(scope="module")
def template():
    return tm_template_rows_v1(ROOT)


def _rows(payload, sheet):
    return [[c.value for c in row] for row in openpyxl.load_workbook(BytesIO(payload))[sheet].iter_rows()]


def test_tm_numbers_are_exactly_the_canonical_coefficients_with_source(example, template):
    _compiled, trial, _index = example
    payload, receipt = build_f36_canonical_workbook_v1(DOCUMENT, trial, template, PINS)
    assert receipt["status"] == "PASS_EXACT_READBACK" and receipt["financial_acceptance"] is False
    assert receipt["tm_values"] == 2 and receipt["canonical_mappings"] == 1
    tm = _rows(payload, "TM")
    assert tm[0][3:] == [f"DURATION:YEAR:{y}|{f}" for y in (2025, 2024) for f in ("value", "unit", "role", "source")]
    row = next(r for r in tm[1:] if r[1] == 1206)
    assert row[3:7] == [1, "MILLION_VND", "CURRENT", f"p64/s1/t1/r1/col2@{trial['mappings'][0]['source_refs'][0]['locator']['page_json_version_id']}"]
    assert row[7:9] == [1, "MILLION_VND"] and row[9] == "COMPARATIVE"
    assert all(r[3] is None for r in tm[1:] if r[1] != 1206)  # no other template row carries a number
    dimensional = _rows(payload, "DIMENSIONAL")
    assert dimensional[0][-2:] == ["ledger_class", "canonical_status"]
    assert [r[-2:] for r in dimensional[1:]] == [["CORE_CANDIDATE", "CANONICAL_MAPPED"]] * 2


def test_withheld_cells_keep_their_named_reason_and_never_reach_tm(example, template):
    _compiled, trial, _index = deepcopy(example)
    ledger = trial["notes_ledger"]
    ledger["retained"] = [{"cell": c, "adapter_disposition": "EXPLICIT_TWO_LANE_PAIR_NOT_PROVEN"} for c in ledger["cells"]]
    trial["mappings"] = []
    payload, receipt = build_f36_canonical_workbook_v1(DOCUMENT, trial, template, PINS)
    assert receipt["tm_values"] == 0 and _rows(payload, "TM")[0] == ["order", "ReportNormId", "name"]
    assert [r[-1] for r in _rows(payload, "DIMENSIONAL")[1:]] == ["RETAINED:EXPLICIT_TWO_LANE_PAIR_NOT_PROVEN"] * 2


def _broken(trial, kind):
    trial = deepcopy(trial)
    mapping = trial["mappings"][0]
    if kind == "duplicate_slot":
        trial["mappings"].append(deepcopy(mapping))
    elif kind == "lane_order":
        for axis in ("values", "period_axis", "source_refs"):
            mapping[axis].reverse()
    elif kind == "not_integer":
        mapping["values"][0]["coefficient"] = 1.5
    elif kind == "foreign_document":
        trial["notes_ledger"]["cells"][0]["document_plan_id"] = "other-plan"
        trial["notes_ledger"]["classifications"][0]["cell"]["document_plan_id"] = "other-plan"
    elif kind == "unaligned_ledger":
        trial["notes_ledger"]["classifications"].reverse()
    return trial


@pytest.mark.parametrize("kind", ["duplicate_slot", "lane_order", "not_integer", "foreign_document", "unaligned_ledger"])
def test_inconsistent_canonical_trials_are_refused(example, template, kind):
    _compiled, trial, _index = example
    with pytest.raises(F36CanonicalWorkbookError):
        build_f36_canonical_workbook_v1(DOCUMENT, _broken(trial, kind), template, PINS)


def test_rnid_outside_the_supplied_template_is_refused_not_dropped(example, template):
    _compiled, trial, _index = example
    with pytest.raises(F36CanonicalWorkbookError, match="outside the supplied template"):
        build_f36_canonical_workbook_v1(DOCUMENT, trial, [r for r in template if r[0] != 1206], PINS)


def test_a_tampered_tm_number_fails_readback(example, template):
    _compiled, trial, _index = example
    payload, _receipt = build_f36_canonical_workbook_v1(DOCUMENT, trial, template, PINS)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    row = next(r for r in workbook["TM"].iter_rows(min_row=2) if r[1].value == 1206)
    row[3].value = 2
    with pytest.raises(NotesWorkbookError):
        verify_f36_canonical_workbook_v1(deterministic_workbook_bytes(workbook, creator="tamper"), DOCUMENT, trial,
                                         template, PINS)
