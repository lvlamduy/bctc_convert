"""Two-lane TM proposal workbook (generalized from the Family-37 workbook tests)."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes
from bctc_ai.export.two_lane_canonical_workbook_v1 import (
    READY,
    TwoLaneCanonicalWorkbookError,
    build_two_lane_canonical_workbook_v1,
    tm_template_rows_v1,
    verify_two_lane_canonical_workbook_v1,
)

FAMILY = "F39 INCOME_TAX"

ROOT = Path(__file__).resolve().parents[2]
DOCUMENT = {"relative_path": "vietstock_bctc/ABB/2025/fixture.pdf", "source_sha256": "c" * 64}


def _ref(row: int, label: str) -> dict:
    return {"label_exact": label, "row_kind": "ITEM", "row_ordinal": row, "money_column_ordinals": [1, 2],
            "locator": {"physical_page": 66, "section_id": "s1", "table_id": "t1",
                        "page_json_version_id": "gfpstorev1:json:" + "b" * 64}}


def _value(coefficient, text, state):
    return {"coefficient": coefficient, "source_text": text, "state": state}


def _trial() -> dict:
    def mapping(role, rnid, values, refs, state="SOURCE_OBSERVED_ROLE_ROW"):
        return {"role": role, "report_norm_id": rnid, "values": values, "source_refs": refs, "state": state,
                "unit": "MILLION_VND"}

    return {"document_ordinal": 1, "status": READY, "mappings": [
        mapping("CUSTOMER_GENERAL", 1224, [_value(118537, "118.537", "RAW_SIGNED_INTEGER"),
                                           _value(-55695, "(55.695)", "RAW_SIGNED_INTEGER")], [_ref(1, "Dự phòng chung")]),
        mapping("OTHER_PROVISION", 1228, [_value(0, "-", "DASH_ZERO"), _value(-54, "(54)", "RAW_SIGNED_INTEGER")],
                [_ref(5, "Hoàn nhập dự phòng khác")]),
        mapping("CUSTOMER_PROVISION", 6031, [_value(143748, None, "EXACT_VISIBLE_GENERAL_AND_SPECIFIC_SUM"),
                                             _value(None, None, "UNOBSERVED_SOURCE_LANE")],
                [_ref(1, "Dự phòng chung"), _ref(2, "Dự phòng cụ thể")], "DECLARED_ROLE_DERIVED_FROM_EXACT_VISIBLE_COMPONENT_SUM"),
    ]}


@pytest.fixture(scope="module")
def template():
    return tm_template_rows_v1(ROOT)


def _tm(payload: bytes) -> dict[int, list]:
    workbook = openpyxl.load_workbook(BytesIO(payload))
    rows = list(workbook["TM"].iter_rows(values_only=True))
    return {row[1]: list(row) for row in rows[1:]}, list(rows[0])


def test_workbook_writes_printed_values_dash_and_derivations_exactly(template) -> None:
    payload, receipt = build_two_lane_canonical_workbook_v1(DOCUMENT, _trial(), template, head="a" * 40, family=FAMILY)
    rows, header = _tm(payload)

    assert header[3:8] == ["CURRENT_PERIOD|value", "CURRENT_PERIOD|state", "CURRENT_PERIOD|unit",
                           "CURRENT_PERIOD|role", "CURRENT_PERIOD|source"]
    assert rows[1224][3] == 118537 and rows[1224][8] == -55695
    assert rows[1228][3] == "-" and rows[1228][4] == "DASH_ZERO"          # a printed dash is never 0
    assert rows[1228][8] == -54
    assert rows[6031][3] == 143748 and rows[6031][4] == "EXACT_VISIBLE_GENERAL_AND_SPECIFIC_SUM"
    assert ";" in rows[6031][7]                                           # both component source rows
    assert rows[6031][8] is None and rows[6031][9] == "UNOBSERVED_SOURCE_LANE"  # nothing zero-filled
    assert rows[1221][3:] == [None] * 10                                  # unmapped root stays empty
    assert receipt["mappings"] == 3 and receipt["tm_cells_written"] == 5
    assert build_two_lane_canonical_workbook_v1(DOCUMENT, _trial(), template, head="a" * 40, family=FAMILY)[0] == payload


@pytest.mark.parametrize("dash", ["-", " - ", "-\n", "--", "_"])
def test_every_printed_dash_form_is_written_as_a_dash(template, dash) -> None:
    trial = _trial()
    trial["mappings"][1]["values"][0] = _value(0, dash, "DASH_ZERO")

    rows, _header = _tm(build_two_lane_canonical_workbook_v1(DOCUMENT, trial, template, head="a" * 40, family=FAMILY)[0])

    assert rows[1228][3] == "-"


def test_readback_rejects_any_changed_cell(template) -> None:
    payload, _receipt = build_two_lane_canonical_workbook_v1(DOCUMENT, _trial(), template, head="a" * 40, family=FAMILY)
    changed = _trial()
    changed["mappings"][0]["values"][0] = _value(118538, "118.538", "RAW_SIGNED_INTEGER")

    with pytest.raises(TwoLaneCanonicalWorkbookError, match="does not read back exactly"):
        verify_two_lane_canonical_workbook_v1(payload, DOCUMENT, changed, template, head="a" * 40, family=FAMILY)


@pytest.mark.parametrize(
    "mutate",
    [
        lambda t: t.update(status="UNRESOLVED_GEMINI_JSON_FAMILY"),
        lambda t: t["mappings"][0].update(report_norm_id=999999),
        lambda t: t["mappings"][1]["values"].__setitem__(0, _value(0, None, "DASH_ZERO")),
        lambda t: t["mappings"][1]["values"].__setitem__(0, _value(0, "0", "DASH_ZERO")),
        lambda t: t["mappings"][2]["values"].__setitem__(1, _value(0, None, "UNOBSERVED_SOURCE_LANE")),
        lambda t: t["mappings"][0]["values"].__setitem__(0, _value(1, "1", "SOME_NEW_STATE")),
        lambda t: t["mappings"].append(dict(t["mappings"][0])),
    ],
)
def test_projection_fails_closed(template, mutate) -> None:
    trial = _trial()
    mutate(trial)

    with pytest.raises(TwoLaneCanonicalWorkbookError):
        build_two_lane_canonical_workbook_v1(DOCUMENT, trial, template, head="a" * 40, family=FAMILY)


def test_producer_incomplete_derived_lane_preserves_other_observed_lane(template) -> None:
    trial = _trial()
    trial["mappings"][2]["values"] = [
        _value(None, None, "DERIVED_INCOMPLETE_DUE_TO_BLANK_SOURCE_CELL"),
        _value(143748, None, "EXACT_VISIBLE_GENERAL_AND_SPECIFIC_SUM"),
    ]
    payload, _receipt = build_two_lane_canonical_workbook_v1(DOCUMENT, trial, template, head="a" * 40, family=FAMILY)
    rows, _header = _tm(payload)
    assert rows[6031][3] is None
    assert rows[6031][4] == "DERIVED_INCOMPLETE_DUE_TO_BLANK_SOURCE_CELL"
    assert rows[6031][8] == 143748
    assert rows[6031][9] == "EXACT_VISIBLE_GENERAL_AND_SPECIFIC_SUM"


def test_producer_provision_reversal_sum_preserves_signed_coefficient(template) -> None:
    trial = _trial()
    mapping = trial["mappings"][0]
    mapping["values"][0] = _value(-12, None, "EXACT_VISIBLE_PROVISION_AND_REVERSAL_SUM")
    mapping["source_refs"] = [_ref(1, "Synthetic provision"), _ref(2, "Synthetic reversal")]
    payload, _receipt = build_two_lane_canonical_workbook_v1(DOCUMENT, trial, template, head="a" * 40, family=FAMILY)
    rows, _header = _tm(payload)
    assert rows[1224][3] == -12
    assert rows[1224][4] == "EXACT_VISIBLE_PROVISION_AND_REVERSAL_SUM"
    assert ";" in rows[1224][7]


@pytest.mark.parametrize(("coefficient", "as_boolean"), [(0, False), (1, True)])
def test_readback_rejects_boolean_equal_to_original_integer(template, coefficient, as_boolean) -> None:
    trial = _trial()
    trial["mappings"][0]["values"][0] = _value(coefficient, str(coefficient), "RAW_SIGNED_INTEGER")
    payload, _receipt = build_two_lane_canonical_workbook_v1(DOCUMENT, trial, template, head="a" * 40, family=FAMILY)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    try:
        row = next(r for r in range(2, workbook["TM"].max_row + 1)
                   if workbook["TM"].cell(r, 2).value == 1224)
        workbook["TM"].cell(row, 4).value = as_boolean
        tampered = deterministic_workbook_bytes(workbook, creator="synthetic boolean tamper")
    finally:
        workbook.close()
    # Same original trial: only actual workbook bytes were changed.
    with pytest.raises(TwoLaneCanonicalWorkbookError, match="does not read back exactly"):
        verify_two_lane_canonical_workbook_v1(tampered, DOCUMENT, trial, template, head="a" * 40, family=FAMILY)
