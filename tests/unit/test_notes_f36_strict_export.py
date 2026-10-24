"""Synthetic strict workbook projection/readback, not corpus or financial acceptance."""

from __future__ import annotations

import json
from copy import deepcopy
from io import BytesIO

import openpyxl
import pytest

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes
from bctc_ai.notes import export_v1 as pilot
from bctc_ai.notes.f36_ledger_projection_v1 import LedgerProjectionError, attach_f36_ledger_v1
from bctc_ai.notes.f36_strict_export_v1 import (
    PRODUCT_LABEL,
    StrictF36WorkbookError,
    build_strict_f36_workbook_v1,
    verify_strict_f36_workbook_v1,
)
from bctc_ai.notes.f36_strict_ledger_v1 import classify_observation_v1
from bctc_ai.source_structure.contracts_v1 import canonical_json_bytes_v1
from test_notes_f36_ledger_projection import REFS
from test_notes_f36_strict_ledger import CLOSED, cell, inventory_row

DOCUMENT = {"document_plan_id": "plan", "relative_path": "synthetic.pdf", "source_sha256": "a" * 64}
TEMPLATE = [(1205, "Tổng"), (1206, "Chi nộp thuế"), (1212, "Chi về tài sản"), (747, "Nhóm 1")]
PINS = {"synthetic_config": "b" * 64}


def _f36(*, row=1, raw="8", value="8", period="DURATION:YEAR:2025", state="VALUE", acceptance="ACCEPTED",
         member="TAX_FEES", rnid=1206, closure=CLOSED):
    original = cell(member, raw, rnid, row=row, value=value, period_key=period, observation_state=state,
                    acceptance=acceptance, hold_reasons=["SOURCE_INCOMPLETE"] if acceptance == "HOLD" else [])
    original.update(period_basis="EXPLICIT_HEADER_YEAR", unit_basis="TABLE", label_exact="Chi nộp thuế")
    original["locator"]["physical_page"] = 3
    inventory = {(original["original_cell"]["page_json_version_id"], 1, 1, row):
                 inventory_row(row, [None, raw, "7"])}
    record = classify_observation_v1(original, rnid, closure, inventory)
    return attach_f36_ledger_v1([original], [record], ledger_input_refs=REFS)[0]


def _build(cells):
    return build_strict_f36_workbook_v1(DOCUMENT, cells, TEMPLATE, PINS)


def _tm(workbook, rnid):
    return next(row for row in workbook["TM"].iter_rows(min_row=2) if row[1].value == rnid)


def test_strict_tm_only_numeric_core_and_complete_dimensional_with_literal_evidence():
    cells = [_f36(), _f36(row=2, member="RENT_EXPENSE", rnid=None, raw="5", value="5"),
             _f36(row=3, rnid=1212, member="ASSET_EXPENSE", closure=None)]
    before = canonical_json_bytes_v1(cells)
    payload, receipt = _build(cells)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    assert receipt["strict_f36_numeric_cells"] == 1 and receipt["dimensional_rows"] == 3
    assert receipt["source_or_ledger_file_authentication_performed"] is False
    assert receipt["canonical_acceptance_or_family_seal_granted"] is False
    assert _tm(workbook, 1206)[3].value == 8 and _tm(workbook, 1212)[3].value is None
    assert _tm(workbook, 1212)[5].value == "ACCEPTED"  # No rewritten acceptance.
    assert _tm(workbook, 1212)[8].value.startswith("ACCEPTED_LOCAL_NOT_ROOT_CLOSED:")
    dimensional = workbook["DIMENSIONAL"]
    assert dimensional.max_row == 4 and dimensional.cell(1, 19).value == "ledger_class"
    assert dimensional.cell(3, 15).value is None and dimensional.cell(3, 11).value == "5"
    assert dimensional.cell(3, 19).value == "ACCEPTED_SOURCE_ONLY"
    assert json.loads(dimensional.cell(2, 20).value) == cells[0]["ledger_evidence"]
    assert workbook["README"].cell(1, 2).value == PRODUCT_LABEL
    assert canonical_json_bytes_v1(cells) == before
    assert verify_strict_f36_workbook_v1(payload, DOCUMENT, cells, TEMPLATE, PINS) == receipt


def test_unique_core_wins_without_dropping_noncore_observation():
    cells = [_f36(row=2, closure=None, raw="99", value="99"), _f36()]
    payload, receipt = _build(cells)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    assert _tm(workbook, 1206)[3].value == 8 and _tm(workbook, 1206)[8].value == "CORE_CANDIDATE"
    assert receipt["dimensional_rows"] == 2 and workbook["DIMENSIONAL"].cell(2, 11).value == "99"


def test_multiple_noncore_states_and_sources_are_explicit_not_a_selected_source():
    cells = [_f36(row=1, closure=None), _f36(row=2, acceptance="HOLD", raw="9", value="9")]
    payload, receipt = _build(cells)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    row = _tm(workbook, 1206)
    assert row[3].value is None and receipt["strict_f36_numeric_cells"] == 0
    assert json.loads(row[4].value) == {"multiplicity": 2, "states": ["VALUE", "VALUE"]}
    assert json.loads(row[5].value) == ["ACCEPTED", "HOLD"]
    sources = json.loads(row[7].value)
    assert [item["dimensional_row"] for item in sources] == [2, 3]
    assert all(f"/r{i}/" in item["source"] for i, item in enumerate(sources, 1))
    assert len(json.loads(row[8].value)) == 2 and workbook["DIMENSIONAL"].max_row == 3


@pytest.mark.parametrize("duplicate", ["core", "exact_noncore"])
def test_duplicate_core_or_exact_noncore_never_silently_selects_one(duplicate):
    first = _f36(closure=CLOSED if duplicate == "core" else None)
    second = _f36(row=2) if duplicate == "core" else deepcopy(first)
    with pytest.raises(StrictF36WorkbookError):
        _build([first, second])


@pytest.mark.parametrize("state,acceptance,value,raw,expected", [
    ("PRINTED_ZERO", "ACCEPTED", "0", "0", 0),
    ("BLANK", "ACCEPTED_BLANK", None, None, None),
    ("DASH", "ACCEPTED_DASH", None, "-", None),
    ("VALUE", "DUPLICATE_IDENTICAL_CONTROL", "8", "8", None),
])
def test_printed_zero_blank_dash_and_control_remain_distinct(state, acceptance, value, raw, expected):
    original = _f36(state=state, acceptance=acceptance, value=value, raw=raw)
    payload, _receipt = _build([original])
    row = _tm(openpyxl.load_workbook(BytesIO(payload)), 1206)
    assert row[3].value == expected and type(row[3].value) is type(expected)
    assert row[4].value == state and row[5].value == acceptance


def test_full_period_key_sort_template_order_and_existing_large_number_policy():
    periods = ["DURATION:MONTHS_ENDED:6:2025-06-30", "INSTANT:2025-06-30", "INSTANT:2024-12-31"]
    cells = [_f36(row=i, period=period, raw=value, value=value)
             for i, (period, value) in enumerate(zip(periods, ["9007199254740993", "2", "3"], strict=True), 1)]
    payload, _receipt = _build(cells)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    header = [c.value for c in workbook["TM"][1]]
    assert [header[i].removesuffix("|value") for i in range(3, len(header), 6)] == [periods[1], periods[0], periods[2]]
    assert [row[1].value for row in workbook["TM"].iter_rows(min_row=2)] == [rnid for rnid, _name in TEMPLATE]
    assert _tm(workbook, 1206)[9].value == "9007199254740993"


def test_other_family_selection_and_values_follow_existing_exporter_behavior():
    foreign = {k: v for k, v in _f36().items() if k not in {"ledger_class", "ledger_evidence"}}
    foreign.update(family="CREDIT_QUALITY", concept="LOANS_TO_CUSTOMERS", axis="CREDIT_QUALITY_GROUP", member="G1",
                   report_norm_id=747, value="9007199254740993")
    default_periods, default_chosen = pilot._tm_cells([foreign])
    payload, receipt = _build([foreign, _f36()])
    workbook = openpyxl.load_workbook(BytesIO(payload))
    assert len(default_periods) == 1 and default_chosen[(747, foreign["period_key"])] == foreign
    assert _tm(workbook, 747)[3].value == "9007199254740993" and _tm(workbook, 747)[8].value is None
    assert receipt["strict_f36_numeric_cells"] == 1


def test_a_valid_attachment_for_a_different_document_is_not_exported_under_this_readme():
    original = _f36()
    with pytest.raises(StrictF36WorkbookError):
        build_strict_f36_workbook_v1({**DOCUMENT, "document_plan_id": "another-plan"},
                                   [original], TEMPLATE, PINS)


def test_core_rnid_must_exist_in_the_supplied_template_without_dropping_it():
    with pytest.raises(StrictF36WorkbookError):
        build_strict_f36_workbook_v1(DOCUMENT, [_f36()], [(1205, "Tổng")], PINS)


@pytest.mark.parametrize("attack", ["missing", "class", "value", "state", "unit", "evidence", "claim"])
def test_missing_or_tampered_ledger_metadata_blocks_every_export(attack):
    original = _f36()
    if attack == "missing":
        del original["ledger_evidence"]
    elif attack == "class":
        original["ledger_class"] = "ACCEPTED_SOURCE_ONLY"
    elif attack == "evidence":
        original["ledger_evidence"]["classification"]["class"] = "ACCEPTED_CORE"
    elif attack == "claim":
        original["ledger_evidence"]["authentication_performed"] = True
    else:
        original[attack] = "changed"
    with pytest.raises(LedgerProjectionError):
        _build([original])


def test_formula_like_labels_and_evidence_remain_literal_text():
    original = _f36(raw="8", value="8")
    # A source label is not a formula. Refresh classification using this complete
    # synthetic original cell; authentication still belongs to the caller.
    base = {k: v for k, v in original.items() if k not in {"ledger_class", "ledger_evidence"}}
    base["label_exact"] = "=1+1"
    rows = {(base["original_cell"]["page_json_version_id"], 1, 1, 1): inventory_row(1, [None, "8", "7"])}
    record = classify_observation_v1(base, 1206, CLOSED, rows)
    projected = attach_f36_ledger_v1([base], [record], ledger_input_refs=REFS)
    payload, receipt = _build(projected)
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    assert workbook["DIMENSIONAL"].cell(2, 16).value == "=1+1"
    assert workbook["DIMENSIONAL"].cell(2, 16).data_type != "f" and receipt["formula_cells"] == 0


@pytest.mark.parametrize("sheet,row,column,value", [
    ("TM", 3, 4, "8"), ("TM", 3, 9, "ACCEPTED_CORE"),
    ("TM", 3, 10, "extra"), ("TM", 6, 2, 1206),
    ("DIMENSIONAL", 2, 20, "{}"), ("DIMENSIONAL", 2, 21, "extra"),
    ("DIMENSIONAL", 3, 1, "extra"), ("README", 1, 2, "ACCEPTED_PRODUCTION"),
])
def test_typed_readback_rejects_changed_values_metadata_or_extra_grid_cells(sheet, row, column, value):
    cells = [_f36()]
    payload, _receipt = _build(cells)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    workbook[sheet].cell(row, column).value = value
    tampered = deterministic_workbook_bytes(workbook, creator="synthetic tamper")
    with pytest.raises(pilot.NotesWorkbookError):
        verify_strict_f36_workbook_v1(tampered, DOCUMENT, cells, TEMPLATE, PINS)
