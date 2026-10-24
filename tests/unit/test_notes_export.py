"""Notes pilot workbook: TM template rows, dimensional view, exact readback and tamper detection."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes
from bctc_ai.notes.export_v1 import (
    NotesWorkbookError,
    build_notes_workbook_v1,
    tm_template_rows_v1,
    verify_notes_workbook_v1,
)

ROOT = Path(__file__).resolve().parents[2]
DOCUMENT = {"document_plan_id": "plan-1", "relative_path": "x.pdf", "source_sha256": "0" * 64}


def cell(rnid, period, value, acceptance="ACCEPTED", member="G1", state="VALUE"):
    return {"family": "CREDIT_QUALITY", "concept": "LOANS_TO_CUSTOMERS", "axis": "CREDIT_QUALITY_GROUP", "member": member,
            "classification_basis": "AS_PRESENTED", "period_key": period, "period_basis": "EXPLICIT_HEADER_DATE",
            "unit": "MILLION_VND", "unit_basis": "HEADER", "observation_state": state, "value": value, "raw": value,
            "acceptance": acceptance, "hold_reasons": [] if acceptance != "HOLD" else ["UNIT_UNRESOLVED"],
            "report_norm_id": rnid, "label_exact": "Nhóm 1",
            "locator": {"physical_page": 3, "section_id": "s1", "table_id": "t1", "row_id": "r1", "column_id": "c1",
                        "page_json_version_id": "v3"}}


@pytest.fixture(scope="module")
def template():
    return tm_template_rows_v1(ROOT)


def test_template_rows_are_pinned_and_complete(template):
    assert len(template) == 1721 and template[0][0] and len({r for r, _ in template}) == 1721


def test_workbook_reads_back_exactly_and_keeps_large_integers_exact(template):
    cells = [cell(747, "INSTANT:2025-06-30", "9007199254740993"), cell(747, "INSTANT:2024-12-31", "800"),
             cell(748, "INSTANT:2025-06-30", None, acceptance="HOLD", member="G2"),
             cell(None, "INSTANT:2025-06-30", "5", member="SAVINGS_TERM")]
    payload, receipt = build_notes_workbook_v1(DOCUMENT, cells, template, {"axes": "x"})
    assert receipt["status"] == "PASS_EXACT_READBACK" and receipt["tm_cells_checked"] == 3 and receipt["dimensional_rows"] == 4
    sheet = openpyxl.load_workbook(io.BytesIO(payload))["TM"]
    header = [c.value for c in sheet[1]]
    assert header[3] == "INSTANT:2025-06-30|value" and header[8] == "INSTANT:2024-12-31|value"
    row = next(r for r in sheet.iter_rows(min_row=2) if r[1].value == 747)
    assert row[3].value == "9007199254740993"  # beyond double precision: exact text


def test_two_accepted_cells_for_one_template_cell_are_refused(template):
    with pytest.raises(NotesWorkbookError):
        build_notes_workbook_v1(DOCUMENT, [cell(747, "INSTANT:2025-06-30", "1"), cell(747, "INSTANT:2025-06-30", "2")],
                                template, {})


def test_tampered_value_is_detected(template):
    cells = [cell(747, "INSTANT:2025-06-30", "900")]
    payload, _ = build_notes_workbook_v1(DOCUMENT, cells, template, {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    row = next(r for r in workbook["TM"].iter_rows(min_row=2) if r[1].value == 747)
    row[3].value = 901
    tampered = deterministic_workbook_bytes(workbook, creator="tamper")
    with pytest.raises(NotesWorkbookError):
        verify_notes_workbook_v1(tampered, DOCUMENT, cells, template, {})


def test_semantic_state_is_visible_on_the_dimensional_sheet_and_read_back(template):
    # Codex 0133: an unproven-meaning observation (5729/5730 lines, no legacy ID) must not look like a proved amount.
    unproven = {**cell(None, "INSTANT:2025-06-30", "5", member="NON_DEDUCTIBLE_EXPENSES"),
                "semantic_state": "BASE_OR_TAX_EFFECT_UNPROVEN"}
    cells = [cell(747, "INSTANT:2025-06-30", "900"), unproven]
    payload, receipt = build_notes_workbook_v1(DOCUMENT, cells, template, {})
    assert receipt["format_version"] == "NOTES_PILOT_WORKBOOK_V2" and receipt["status"] == "PASS_EXACT_READBACK"
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    sheet = workbook["DIMENSIONAL"]
    assert sheet.cell(1, 18).value == "semantic_state"
    assert [sheet.cell(row, 18).value for row in (2, 3)] == [None, "BASE_OR_TAX_EFFECT_UNPROVEN"]
    assert all(r[1].value != "NON_DEDUCTIBLE_EXPENSES" for r in workbook["TM"].iter_rows(min_row=2))  # never a template row
    sheet.cell(3, 18).value = None
    with pytest.raises(NotesWorkbookError):
        verify_notes_workbook_v1(deterministic_workbook_bytes(workbook, creator="tamper"), DOCUMENT, cells, template, {})


def _tamper(payload, edit):
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    edit(workbook)
    return deterministic_workbook_bytes(workbook, creator="tamper")


def _set(sheet, row, column, value):
    def edit(workbook):
        workbook[sheet].cell(row, column).value = value
    return edit


def test_every_emitted_cell_and_axis_is_read_back_with_its_type(template):
    # Codex 0222: the verifier compares the complete projection (README, TM header/order/name/periods, all 18
    # DIMENSIONAL columns) with exact types, and rejects cells outside it.
    cells = [cell(747, "INSTANT:2025-06-30", "900"), cell(747, "INSTANT:2024-12-31", "800"),
             {**cell(748, "INSTANT:2025-06-30", None, acceptance="HOLD", member="G2", state="BLANK"), "raw": None}]
    pins = {"axes": "x"}
    payload, receipt = build_notes_workbook_v1(DOCUMENT, cells, template, pins)
    assert receipt["readback_scope"] == "EVERY_EMITTED_CELL_TYPED_NO_EXTRA_CELLS" and receipt["status"] == "PASS_EXACT_READBACK"
    assert verify_notes_workbook_v1(payload, DOCUMENT, cells, template, pins)["grid_positions_compared"] == receipt["grid_positions_compared"]
    tm = openpyxl.load_workbook(io.BytesIO(payload))["TM"]
    row747 = next(r[0].row for r in tm.iter_rows(min_row=2) if r[1].value == 747)
    tampers = {
        "dimensional observation_state BLANK->DASH": _set("DIMENSIONAL", 4, 10, "DASH"),
        "dimensional raw None->'-'": _set("DIMENSIONAL", 4, 12, "-"),
        "dimensional family": _set("DIMENSIONAL", 2, 1, "OPERATING_EXPENSE"),
        "dimensional axis": _set("DIMENSIONAL", 2, 3, "OTHER_AXIS"),
        "dimensional classification_basis": _set("DIMENSIONAL", 2, 5, "TT31_2024"),
        "dimensional period_basis": _set("DIMENSIONAL", 2, 7, "ALIGNED_TO_PRIMARY_BALANCE_CURRENT_LANE"),
        "dimensional unit_basis": _set("DIMENSIONAL", 2, 9, "TABLE"),
        "dimensional hold_reasons": _set("DIMENSIONAL", 4, 14, ""),
        "dimensional report_norm_id": _set("DIMENSIONAL", 2, 15, 746),
        "dimensional report_norm_id as text": _set("DIMENSIONAL", 2, 15, "747"),
        "dimensional label_exact": _set("DIMENSIONAL", 2, 16, "Nhóm 2"),
        "dimensional header": _set("DIMENSIONAL", 1, 11, "amount"),
        "dimensional extra row": _set("DIMENSIONAL", 5, 1, "CREDIT_QUALITY"),
        "dimensional extra column": _set("DIMENSIONAL", 2, 19, "x"),
        "tm period header": _set("TM", 1, 4, "INSTANT:2025-03-31|value"),
        "tm order": _set("TM", row747, 1, 9999),
        "tm name": _set("TM", row747, 3, "renamed"),
        "tm value type int->text": _set("TM", row747, 4, "900"),
        "tm extra column": _set("TM", row747, 14, "x"),
        "tm extra row": _set("TM", len(template) + 2, 2, 1),
        "readme relative_path": _set("README", 4, 2, "y.pdf"),
        "readme source_sha256": _set("README", 5, 2, "1" * 64),
        "readme families": _set("README", 6, 2, "CREDIT_QUALITY,INCOME_TAX"),
        "readme pins": _set("README", 7, 2, "[]"),
    }
    for name, edit in tampers.items():
        with pytest.raises(NotesWorkbookError):
            verify_notes_workbook_v1(_tamper(payload, edit), DOCUMENT, cells, template, pins)
        assert name  # every tamper above is a distinct, previously unchecked field or axis
    with pytest.raises(NotesWorkbookError):  # the same bytes against other pins are not the same workbook
        verify_notes_workbook_v1(payload, DOCUMENT, cells, template, {"axes": "y"})
