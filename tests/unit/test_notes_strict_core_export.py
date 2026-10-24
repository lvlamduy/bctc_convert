"""Opt-in strict-core notes workbook (Codex 0357 / 0359): F36 numbers only for strict-core eligible cells."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes
from bctc_ai.notes.export_v1 import NotesWorkbookError, build_notes_workbook_v1, tm_template_rows_v1
from bctc_ai.notes.strict_core_export_v1 import (
    build_notes_strict_core_workbook_v1,
    verify_notes_strict_core_workbook_v1,
)

ROOT = Path(__file__).resolve().parents[2]
DOCUMENT = {"document_plan_id": "plan-69", "relative_path": "x.pdf", "source_sha256": "0" * 64}
PERIOD = "DURATION:YEAR:2025"


def opex(rnid, member, value, acceptance="ACCEPTED", state="VALUE", ledger_class="CORE_CANDIDATE", row="r1"):
    """An OPERATING_EXPENSE engine cell plus the root projection's additive ledger keys."""

    return {"family": "OPERATING_EXPENSE", "concept": "OPERATING_EXPENSE_NOTE", "axis": "OPERATING_EXPENSE_LINE",
            "member": member, "classification_basis": "AS_PRESENTED", "period_key": PERIOD, "period_basis": "EXPLICIT_HEADER_DATE",
            "unit": "MILLION_VND", "unit_basis": "HEADER", "observation_state": state, "value": value, "raw": value,
            "acceptance": acceptance, "hold_reasons": ["LINE_NOT_PROVEN_BY_ITS_OWN_EQUATION_OR_ANCHOR"] if acceptance == "HOLD" else [],
            "report_norm_id": rnid, "label_exact": member,
            "locator": {"physical_page": 64, "section_id": "s1", "table_id": "t1", "row_id": row, "column_id": "c2",
                        "page_json_version_id": "v64"},
            "ledger_class": ledger_class, "ledger_evidence": {"class": ledger_class, "summary_ref": {"sha256": "a" * 64}}}


def credit(value="800", acceptance="ACCEPTED"):
    return {"family": "CREDIT_QUALITY", "concept": "LOANS_TO_CUSTOMERS", "axis": "CREDIT_QUALITY_GROUP", "member": "G1",
            "classification_basis": "AS_PRESENTED", "period_key": PERIOD, "period_basis": "EXPLICIT_HEADER_DATE",
            "unit": "MILLION_VND", "unit_basis": "HEADER", "observation_state": "VALUE", "value": value, "raw": value,
            "acceptance": acceptance, "hold_reasons": [], "report_norm_id": 747, "label_exact": "Nhóm 1",
            "locator": {"physical_page": 3, "section_id": "s1", "table_id": "t1", "row_id": "r1", "column_id": "c1",
                        "page_json_version_id": "v3"}}


@pytest.fixture(scope="module")
def template():
    return tm_template_rows_v1(ROOT)


def _cells():
    return [opex(1206, "TAX_FEES", "110573"),
            opex(1207, "EMPLOYEE_EXPENSE", "900", ledger_class="ACCEPTED_LOCAL_NOT_ROOT_CLOSED:UNRESOLVED:PARENT_EDGE_DIFFERS_FROM_ORIGINAL:r4", row="r2"),
            opex(1208, "SALARY_ALLOWANCE", None, acceptance="ACCEPTED_BLANK", state="BLANK", ledger_class="ACCEPTED_PRINTED_BLANK_OR_DASH", row="r3"),
            opex(1209, "PAYROLL_CONTRIBUTIONS", "7", acceptance="HOLD", ledger_class="HOLD:LINE_NOT_PROVEN_BY_ITS_OWN_EQUATION_OR_ANCHOR", row="r4"),
            opex(None, "RENT_EXPENSE", "761240", ledger_class="ACCEPTED_SOURCE_ONLY", row="r5"),
            credit()]


ELIGIBLE = [True, False, False, False, False, False]


def _tm_row(payload, rnid):
    sheet = openpyxl.load_workbook(io.BytesIO(payload))["TM"]
    return next([c.value for c in row] for row in sheet.iter_rows(min_row=2) if row[1].value == rnid)


def test_the_default_export_is_byte_identical_with_or_without_ledger_keys(template):
    plain = [{k: v for k, v in c.items() if k not in {"ledger_class", "ledger_evidence"}} for c in _cells()]
    assert build_notes_workbook_v1(DOCUMENT, _cells(), template, {})[0] == build_notes_workbook_v1(DOCUMENT, plain, template, {})[0]


def test_only_a_strict_core_eligible_f36_cell_carries_a_number(template):
    payload, receipt = build_notes_strict_core_workbook_v1(DOCUMENT, _cells(), ELIGIBLE, template, {"ledger": "x"})
    assert receipt["status"] == "PASS_EXACT_READBACK" and receipt["financial_acceptance"] is False
    assert _tm_row(payload, 1206)[3:9] == [110573, "VALUE", "ACCEPTED", "MILLION_VND", "p64/s1/t1/r1/c2@v64", "CORE_CANDIDATE"]
    local = _tm_row(payload, 1207)
    assert local[3] is None and local[5] == "ACCEPTED" and local[8].startswith("ACCEPTED_LOCAL_NOT_ROOT_CLOSED:")
    blank = _tm_row(payload, 1208)
    assert blank[3] is None and blank[3] != 0 and blank[4:6] == ["BLANK", "ACCEPTED_BLANK"]  # printed state, never 0
    assert _tm_row(payload, 1209)[3] is None
    assert _tm_row(payload, 747)[3:9] == [800, "VALUE", "ACCEPTED", "MILLION_VND", "p3/s1/t1/r1/c1@v3", "NOT_LEDGERED"]
    assert receipt["tm_numbers_emitted"] == 2  # the core F36 cell and the unledgered family's default number


def test_dimensional_keeps_every_observation_with_its_class_and_evidence(template):
    payload, receipt = build_notes_strict_core_workbook_v1(DOCUMENT, _cells(), ELIGIBLE, template, {})
    rows = [[c.value for c in r] for r in openpyxl.load_workbook(io.BytesIO(payload))["DIMENSIONAL"].iter_rows()]
    assert rows[0][-3:] == ["ledger_class", "strict_core", "ledger_evidence"] and len(rows) == 7 == receipt["dimensional_rows"] + 1
    rent = next(r for r in rows if r[3] == "RENT_EXPENSE")
    assert rent[-3:-1] == ["ACCEPTED_SOURCE_ONLY", "NOT_ELIGIBLE"] and rent[14] is None  # source-only: no invented RNID
    assert rows[1][-1] == '{"class":"CORE_CANDIDATE","summary_ref":{"sha256":"' + "a" * 64 + '"}}'
    assert rows[-1][-3:] == ["NOT_LEDGERED", "NOT_ELIGIBLE", None]


_DROP = object()


def _broken(index, **change):
    cells = _cells()
    cells[index] = {k: v for k, v in {**cells[index], **change}.items() if v is not _DROP}
    return cells


@pytest.mark.parametrize("cells,eligible", [
    (_broken(0, ledger_class=_DROP), ELIGIBLE),
    (_broken(0, ledger_evidence=None), ELIGIBLE),
    (_broken(1, ledger_class=""), ELIGIBLE),
    (_broken(5, ledger_class="CORE_CANDIDATE"), ELIGIBLE),
    (_cells(), ELIGIBLE[:-1]),
    (_cells(), [1, 0, 0, 0, 0, 0]),
    (_cells(), [True, True, False, False, False, False]),
    (_broken(0, acceptance="HOLD"), ELIGIBLE),
])
def test_missing_or_misplaced_ledger_metadata_and_wrong_eligibility_are_refused(template, cells, eligible):
    with pytest.raises(NotesWorkbookError):
        build_notes_strict_core_workbook_v1(DOCUMENT, cells, eligible, template, {})


def test_a_number_smuggled_into_a_non_core_f36_position_fails_readback(template):
    cells = _cells()
    payload, _receipt = build_notes_strict_core_workbook_v1(DOCUMENT, cells, ELIGIBLE, template, {})
    workbook = openpyxl.load_workbook(io.BytesIO(payload))
    sheet = workbook["TM"]
    row = next(r for r in sheet.iter_rows(min_row=2) if r[1].value == 1207)
    row[3].value = 900
    with pytest.raises(NotesWorkbookError):
        verify_notes_strict_core_workbook_v1(deterministic_workbook_bytes(workbook, creator="tamper"), DOCUMENT, cells,
                                             ELIGIBLE, template, {})


def test_the_duplicate_report_norm_period_refusal_is_kept(template):
    cells = _cells() + [opex(1206, "TAX_FEES", "1", row="r9")]
    with pytest.raises(NotesWorkbookError):
        build_notes_strict_core_workbook_v1(DOCUMENT, cells, ELIGIBLE + [False], template, {})
