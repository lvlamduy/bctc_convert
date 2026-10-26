"""Synthetic publication wiring and serialized workbook rejection controls.

The publication fixture replaces only the upstream store-admission callback.
It does not authenticate a real store, PDF, family or source-frontier receipt.
Projection, the pinned repository template, file authentication, workbook
render/readback, output guards and exclusive publication remain production code.
The separate bridge test module owns the real synthetic store/source E2E.
"""

from __future__ import annotations

import copy
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import test_gemini_tm_multitable_bridge_v1 as fixtures

from bctc_ai.export import gemini_projected_tm_workbook_v1 as renderer
from bctc_ai.export import gemini_tm_multitable_bridge_v1 as bridge


def _publication_wiring_case(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> tuple[dict[str, Any], dict[str, Any], list[dict[str, Any]]]:
    """Admit only explicitly synthetic inputs; never bypass downstream guards."""
    inputs = tmp_path / "synthetic-input-authorities"
    inputs.mkdir()
    roots = {}
    for key in ("source_pdf_root", "corpus_artifact_root", "producer_root",
                "effective_page_artifact_root"):
        root = inputs / key
        root.mkdir()
        roots[key] = str(root)

    def synthetic_ref(name: str, payload: bytes) -> dict[str, Any]:
        path = inputs / name
        path.write_bytes(payload)
        return bridge.file_ref_v1(path)

    template = Path(__file__).resolve().parents[2] / "template" / "Bank_TM_ReportNormId.v2.xlsx"
    bundle = fixtures._bundle()
    bundle["refs"] = {
        "store_ref": synthetic_ref("NOT-A-REAL-STORE.sqlite3", b"synthetic admission seam only"),
        "corpus_index_ref": synthetic_ref("synthetic-corpus.json", b'{"synthetic_only":true}\n'),
        "source_page_database_ref": synthetic_ref("NOT-A-REAL-PAGE-STORE.sqlite3", b"synthetic page seam only"),
        "template_ref": bridge.file_ref_v1(template),
    }
    bundle["producer_files"] = [synthetic_ref("synthetic-producer.txt", b"not an executed producer\n")]
    bundle["pdf_ref"] = synthetic_ref("NOT-A-REAL-PDF.bin", b"synthetic bytes, not PDF evidence\n")
    bundle["run_metadata"]["admission_scope"] = "SYNTHETIC_PUBLICATION_WIRING_ONLY_NOT_STORE_AUTHENTICATION"
    exports = tmp_path / "exports"
    exports.mkdir()
    request = {
        **roots,
        **bundle["refs"],
        "family_run_id": bundle["run_metadata"]["family_run_id"],
        "producer_refs": bundle["producer_files"],
        "document_ordinal": bundle["document"]["document_ordinal"],
        "expected_family_id": bundle["sweep"]["family_id"],
        "expected_run_kind": "EXPERIMENTAL",
        "require_current_selection": False,
        "output_dir": str(exports / "new-document"),
    }
    admissions: list[dict[str, Any]] = []

    def admit_synthetic_only(value: dict[str, Any]) -> dict[str, Any]:
        admissions.append(copy.deepcopy(value))
        return copy.deepcopy(bundle)

    monkeypatch.setattr(bridge, "authenticate_multitable_store_run_v1", admit_synthetic_only)
    return request, bundle, admissions


def _files_below(root: Path) -> dict[str, bytes]:
    return {str(path.relative_to(root)): path.read_bytes()
            for path in root.rglob("*") if path.is_file()}


def test_publication_wiring_links_exact_workbook_receipt_and_provenance(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    request, bundle, admissions = _publication_wiring_case(tmp_path, monkeypatch)
    before = _files_below(tmp_path / "synthetic-input-authorities")
    result = bridge.export_tm_multitable_store_document_v1(request=request)
    output = Path(request["output_dir"])
    assert set(path.name for path in output.iterdir()) == {"tm.xlsx", "receipt.json", "receipt.sha256"}
    assert len(admissions) == 1 and admissions[0] == request
    assert before == _files_below(tmp_path / "synthetic-input-authorities")

    payload = (output / "tm.xlsx").read_bytes()
    receipt_bytes = (output / "receipt.json").read_bytes()
    receipt = bridge.strict_json_v1(receipt_bytes)
    assert result["xlsx_sha256"] == sha256(payload).hexdigest() == receipt["xlsx"]["sha256"]
    assert receipt["xlsx"]["size_bytes"] == len(payload)
    assert receipt["xlsx"]["path"] == str(output / "tm.xlsx")
    assert result["receipt_sha256"] == sha256(receipt_bytes).hexdigest()
    assert (output / "receipt.sha256").read_bytes() == (
        result["receipt_sha256"] + "  receipt.json\n"
    ).encode("ascii")
    material = {key: value for key, value in receipt.items() if key != "receipt_material_sha256"}
    assert receipt["receipt_material_sha256"] == sha256(bridge.json_bytes(material)).hexdigest()
    assert receipt_bytes == bridge.json_bytes(receipt)
    assert receipt["document"] == bundle["document"]
    assert receipt["provenance"]["trial"] == bundle["trial"]
    assert receipt["provenance"]["input_refs"] == bundle["refs"]
    assert receipt["provenance"]["source_pdf_ref"] == bundle["pdf_ref"]
    assert receipt["provenance"]["schema_binding"] == bundle["sweep"]["specs"]["schema_binding"]
    assert receipt["provenance"]["run_metadata"]["admission_scope"] == (
        "SYNTHETIC_PUBLICATION_WIRING_ONLY_NOT_STORE_AUTHENTICATION"
    )
    assert receipt["label"] == renderer.LABEL == result["label"]
    assert receipt["financial_acceptance"] == "NOT_CLAIMED"
    assert receipt["store_modified"] is False
    assert receipt["xlsx"]["typed_readback"] == "PASS"
    records = bridge.project_multitable_document_v1(bundle)
    assert receipt["projection_count"] == len(records)
    assert receipt["provenance"]["projection_sha256"] == sha256(bridge.json_bytes(records)).hexdigest()
    template = renderer.template_rows_v1(bridge.pinned_bytes_v1(bundle["refs"]["template_ref"]))
    expected = renderer.projected_sheet_rows_v1(
        document=bundle["document"], records=records, template=template, provenance=receipt["provenance"],
    )
    renderer.verify_projected_workbook_v1(payload, expected)


@pytest.mark.parametrize("existing_kind", ["directory", "file"])
def test_publication_wiring_preserves_existing_output_and_every_old_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, existing_kind: str,
) -> None:
    request, _, admissions = _publication_wiring_case(tmp_path, monkeypatch)
    output = Path(request["output_dir"])
    if existing_kind == "directory":
        output.mkdir()
        for name in ("tm.xlsx", "receipt.json", "receipt.sha256", "owner-notes.txt"):
            (output / name).write_bytes(("ORIGINAL " + name).encode("ascii"))
    else:
        output.write_bytes(b"ORIGINAL file at requested directory path")
    before = _files_below(tmp_path)
    with pytest.raises(ValueError, match="output directory must be new and absolute"):
        bridge.export_tm_multitable_store_document_v1(request=request)
    assert len(admissions) == 1
    assert _files_below(tmp_path) == before


def test_publication_wiring_loses_creation_race_without_overwriting_new_owner(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    request, _, _ = _publication_wiring_case(tmp_path, monkeypatch)
    output = Path(request["output_dir"])
    before = _files_below(tmp_path / "synthetic-input-authorities")
    original_mkdir = Path.mkdir

    def concurrent_owner(path: Path, *args: Any, **kwargs: Any) -> None:
        if path == output:
            original_mkdir(path)
            (path / "tm.xlsx").write_bytes(b"CONCURRENT OWNER workbook")
            (path / "receipt.json").write_bytes(b"CONCURRENT OWNER receipt")
        original_mkdir(path, *args, **kwargs)

    # Only simulate the competing filesystem writer after the absent-path
    # check. The real exclusive mkdir, path guards and readback still execute.
    monkeypatch.setattr(Path, "mkdir", concurrent_owner)
    with pytest.raises(FileExistsError):
        bridge.export_tm_multitable_store_document_v1(request=request)
    assert _files_below(output) == {
        "tm.xlsx": b"CONCURRENT OWNER workbook", "receipt.json": b"CONCURRENT OWNER receipt",
    }
    assert before == _files_below(tmp_path / "synthetic-input-authorities")


@pytest.mark.parametrize("root_key", [
    "source_pdf_root", "corpus_artifact_root", "producer_root", "effective_page_artifact_root",
    "authority_file_parent",
])
def test_publication_wiring_cannot_create_output_inside_input_authority(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, root_key: str,
) -> None:
    request, bundle, _ = _publication_wiring_case(tmp_path, monkeypatch)
    parent = (Path(bundle["refs"]["store_ref"]["path"]).parent if root_key == "authority_file_parent"
              else Path(request[root_key]))
    request["output_dir"] = str(parent / "forbidden-new-output")
    before = _files_below(tmp_path)
    with pytest.raises(ValueError, match="output overlaps input authority"):
        bridge.export_tm_multitable_store_document_v1(request=request)
    assert not Path(request["output_dir"]).exists()
    assert _files_below(tmp_path) == before


@pytest.mark.parametrize("git_marker", ["directory", "worktree_file"])
def test_publication_wiring_cannot_create_output_in_git_worktree(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, git_marker: str,
) -> None:
    request, _, _ = _publication_wiring_case(tmp_path, monkeypatch)
    worktree = tmp_path / "synthetic-git-worktree"
    worktree.mkdir()
    marker = worktree / ".git"
    if git_marker == "directory":
        marker.mkdir()
    else:
        marker.write_bytes(b"gitdir: synthetic marker only\n")
    (worktree / "owner.txt").write_bytes(b"keep this unchanged")
    request["output_dir"] = str(worktree / "new-output")
    before = _files_below(tmp_path)
    with pytest.raises(ValueError, match="Git worktree"):
        bridge.export_tm_multitable_store_document_v1(request=request)
    assert not Path(request["output_dir"]).exists()
    assert _files_below(tmp_path) == before


def test_publication_wiring_cannot_publish_through_symlink_parent(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    request, _, _ = _publication_wiring_case(tmp_path, monkeypatch)
    target = tmp_path / "real-owner-directory"
    target.mkdir()
    (target / "owner.txt").write_bytes(b"original owner bytes")
    link = tmp_path / "linked-output-parent"
    try:
        link.symlink_to(target, target_is_directory=True)
    except (OSError, NotImplementedError):
        pytest.skip("platform does not permit directory symlink creation")
    request["output_dir"] = str(link / "new-output")
    before = _files_below(target)
    with pytest.raises(ValueError, match="symlink/junction"):
        bridge.export_tm_multitable_store_document_v1(request=request)
    assert not (target / "new-output").exists()
    assert _files_below(target) == before


def _readback_case() -> tuple[bytes, dict[str, list[list[Any]]]]:
    bundle = fixtures._bundle()
    kwargs = {
        "document": bundle["document"], "records": bridge.project_multitable_document_v1(bundle),
        "template": [(1248, "family root"), (1249, "cash"), (1250, "=literal schema name")],
        "provenance": {"trial": bundle["trial"], "scope": "SYNTHETIC_READBACK_ONLY"},
    }
    expected = renderer.projected_sheet_rows_v1(**kwargs)
    payload, _ = renderer.render_projected_tm_workbook_v1(**kwargs)
    return payload, expected


def _replace_literal(cell: Any, value: Any) -> None:
    cell.value = value
    if type(value) is str:
        cell.data_type = "s"


def _swap_rows(sheet: Any, first: int, second: int) -> None:
    left = [cell.value for cell in sheet[first]]
    right = [cell.value for cell in sheet[second]]
    for column, (first_value, second_value) in enumerate(zip(left, right, strict=True), 1):
        _replace_literal(sheet.cell(first, column), second_value)
        _replace_literal(sheet.cell(second, column), first_value)


@pytest.mark.parametrize("mutation", [
    "sheet_order", "sheet_name", "tm_row_order", "provenance_row_order", "status_row_order",
    "schema_rnid", "schema_name", "provenance_json", "provenance_chunk_axis", "status_acceptance",
    "numeric_as_text", "formula", "extra_value_column", "extra_value_row",
])
def test_serialized_readback_rejects_schema_every_sheet_and_provenance_tamper(mutation: str) -> None:
    payload, expected = _readback_case()
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    try:
        if mutation == "sheet_order":
            workbook.move_sheet("Status", offset=-2)
        elif mutation == "sheet_name":
            workbook["Provenance"].title = "UnlinkedProvenance"
        elif mutation == "tm_row_order":
            _swap_rows(workbook["TM"], 2, 3)
        elif mutation == "provenance_row_order":
            _swap_rows(workbook["Provenance"], 2, 3)
        elif mutation == "status_row_order":
            _swap_rows(workbook["Status"], 1, 2)
        elif mutation == "schema_rnid":
            workbook["TM"].cell(3, 1).value = 1250
        elif mutation == "schema_name":
            _replace_literal(workbook["TM"].cell(3, 2), "different schema label")
        elif mutation == "provenance_json":
            _replace_literal(workbook["Provenance"].cell(2, 6), '{"document_id":"tampered"}')
        elif mutation == "provenance_chunk_axis":
            workbook["Provenance"].cell(2, 4).value = 999
        elif mutation == "status_acceptance":
            _replace_literal(workbook["Status"].cell(4, 2), "FINANCIALLY_ACCEPTED")
        elif mutation == "numeric_as_text":
            column = next(index for index, value in enumerate(expected["TM"][0], 1)
                          if "2025-12-31" in str(value) and str(value).endswith("::Value"))
            assert type(workbook["TM"].cell(3, column).value) is int
            _replace_literal(workbook["TM"].cell(3, column), "-123")
        elif mutation == "formula":
            workbook["TM"].cell(3, 3).value = "=1+1"
        elif mutation == "extra_value_column":
            _replace_literal(workbook["Provenance"].cell(2, 7), "unbound metadata")
        else:
            _replace_literal(workbook["Status"].cell(7, 1), "unbound extra status")
        output = BytesIO()
        workbook.save(output)
    finally:
        workbook.close()
    with pytest.raises(ValueError, match="(axis drift|typed workbook value drift|contains a formula)"):
        renderer.verify_projected_workbook_v1(output.getvalue(), expected)
