"""Synthetic product controls; no restored corpus/PDF/SQLite inputs are used."""

from __future__ import annotations

import copy
import json
import sqlite3
from hashlib import sha256
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from bctc_ai.export import gemini_projected_tm_workbook_v1 as renderer
from bctc_ai.export import gemini_tm_multitable_bridge_v1 as bridge


def _bundle() -> dict:
    document = {"document_id": "gfpstorev1:document:" + "a" * 64,
                "document_plan_id": "gjfpdocv1:" + "b" * 64,
                "document_ordinal": 3, "source_ordinal": 3,
                "source_logical_name": "synthetic.pdf", "source_sha256": "c" * 64,
                "relative_path": "synthetic.pdf", "status": "READY_FOR_SCHEMA_MAPPING_REVIEW_PROPOSAL_ONLY"}
    version = "gfpstorev1:json:" + "d" * 64
    axis = {key: document[key] for key in ("document_id", "document_ordinal", "source_logical_name", "source_sha256")}
    axis.update(page_json_version_id=version, physical_page=7, selected_page_ordinal=1)
    locator = {**axis, "section_id": "s1", "table_id": "t1", "component_roles": ["CASH"]}
    row = {"label_exact": "Tiền mặt", "hierarchy_path_exact": ["Tiền mặt"],
           "row_kind": "ITEM", "values_exact": ["(123)", "-"]}
    source = {"locator": locator, "money_column_ordinals": [1, 2], "row_id": "r1",
              "row_ordinal": 1, **{key: row[key] for key in ("label_exact", "hierarchy_path_exact", "row_kind")}}
    lanes = {"complete": True, "money_column_ordinals": [1, 2],
             "lane_keys": [["SEMANTIC_ALIAS", "CURRENT_PERIOD"], ["SEMANTIC_ALIAS", "COMPARATIVE_PERIOD"]],
             "source_lane_keys": [["DATE", "2025-12-31"], ["DATE", "2024-12-31"]]}
    receipt = {"region": locator, "lane_axis": lanes,
               "unit_axis": {"complete": True, "canonical_unit": "VND_MILLION", "source_exact": "Triệu đồng"}}
    mapping = {"item_mapping_id": "gjmthfmv1:item:" + "e" * 64,
               "role": "CASH", "report_norm_id": 1249, "row_id": "r1",
               "source_refs": [source], "state": "SOURCE_OBSERVED_ROW", "unit": "VND_MILLION",
               "values": [{"coefficient": -123, "source_text": "(123)", "state": "RAW_SIGNED_INTEGER"},
                          {"coefficient": 0, "source_text": "-", "state": "DASH_ZERO"}]}
    candidate = {"candidate_id": "candidate:test", "mappings": [mapping],
                 "closure_receipt": {"table_receipts": [receipt], "source_only_unmapped_rows": [{"synthetic": "retained"}]}}
    trial = {"document_ordinal": 3, "source_logical_name": "synthetic.pdf", "source_sha256": "c" * 64,
             "status": document["status"], "selected_candidate_id": candidate["candidate_id"],
             "mappings": [mapping], "candidates": [candidate], "reasons": []}
    table = {"columns": [{"header_path_exact": ["31/12/2025", "Triệu đồng"]},
                          {"header_path_exact": ["31/12/2024", "Triệu đồng"]}], "rows": [row]}
    schema = {"role_bindings": [{"role": "CASH", "report_norm_id": 1249}],
              "family_root_report_norm_id": 1248, "schema_period_type": "SNAPSHOT"}
    return {"document": document, "trial": trial,
            "sweep": {"family_id": "CASH_EQUIVALENTS", "sweep_id": "synthetic:sweep",
                      "specs": {"schema_binding": {"value": schema}}},
            "run_metadata": {"family_run_id": "synthetic:run"},
            "page_axis": [axis], "pages": {version: {"page_json": {"sections": [{"tables": [table]}]}}}}


def _mapping(bundle: dict) -> dict:
    return bundle["trial"]["mappings"][0]


def _receipt(bundle: dict) -> dict:
    return bundle["trial"]["candidates"][0]["closure_receipt"]["table_receipts"][0]


def _raw(bundle: dict) -> dict:
    return next(iter(bundle["pages"].values()))["page_json"]["sections"][0]["tables"][0]["rows"][0]


def _render(bundle: dict) -> tuple[bytes, dict]:
    return renderer.render_projected_tm_workbook_v1(
        document=bundle["document"], records=bridge.project_multitable_document_v1(bundle),
        template=[(1248, "Tổng"), (1249, "Tiền mặt"), (1250, "=literal name")],
        provenance={"trial": bundle["trial"], "synthetic_only": True},
    )


def test_raw_integer_dash_period_unit_and_global_ordinal_survive_projection() -> None:
    bundle = _bundle()
    records = bridge.project_multitable_document_v1(bundle)
    assert len(records) == 2
    assert records[0]["cell"]["coefficient"] == -123
    assert records[0]["period"] == {"basis": "STOCK_AS_OF", "as_of": "2025-12-31"}
    assert records[1]["period"]["as_of"] == "2024-12-31"  # Allowed comparative column.
    assert records[0]["global_document_ordinal"] == 3
    assert records[0]["source_refs"] == _mapping(bundle)["source_refs"]
    assert records[0]["unit"] == "VND_MILLION"  # No magnitude conversion.


def test_full_flow_range_is_supported_without_guessing_year_start() -> None:
    bundle = _bundle()
    bundle["sweep"]["specs"]["schema_binding"]["value"]["schema_period_type"] = "DURATION"
    _receipt(bundle)["lane_axis"]["source_lane_keys"] = [
        ["DURATION_DATE_RANGE", "2025-04-01", "2025-06-30"],
        ["DURATION_DATE_RANGE", "2024-04-01", "2024-06-30"],
    ]
    records = bridge.project_multitable_document_v1(bundle)
    assert records[0]["period"] == {"basis": "FLOW_RANGE", "start": "2025-04-01", "end": "2025-06-30"}


@pytest.mark.parametrize("key,kind", [(["BARE_YEAR", "2025"], "SNAPSHOT"),
    (["SEMANTIC_ALIAS", "CURRENT_PERIOD"], "SNAPSHOT"),
    (["DURATION_END_DATE", "2025-06-30"], "DURATION"),
    (["DURATION_BARE_YEAR", "2025"], "DURATION")])
def test_incomplete_period_proof_is_rejected(key: list, kind: str) -> None:
    with pytest.raises(ValueError, match="full source-proven"):
        bridge.period_from_source_lane_v1(key, kind)


@pytest.mark.parametrize("mutation", ["source", "source_ordinal", "page", "row", "unit", "unit_incomplete", "column", "role", "rnid", "coefficient", "semantic_role", "derived"])
def test_each_source_binding_tamper_stops_the_document(mutation: str) -> None:
    bundle = _bundle()
    mapping = _mapping(bundle)
    source = mapping["source_refs"][0]
    if mutation == "source":
        source["locator"]["source_sha256"] = "f" * 64
    elif mutation == "source_ordinal":
        source["locator"]["document_ordinal"] = 1
    elif mutation == "page":
        source["locator"]["physical_page"] = 8
    elif mutation == "row":
        _raw(bundle)["label_exact"] = "Different row"
    elif mutation == "unit":
        mapping["unit"] = "VND"
    elif mutation == "unit_incomplete":
        _receipt(bundle)["unit_axis"]["complete"] = False
    elif mutation == "column":
        source["money_column_ordinals"] = [2, 1]
    elif mutation == "role":
        mapping["role"] = "OTHER_FAMILY"
    elif mutation == "rnid":
        mapping["report_norm_id"] = 999999
    elif mutation == "coefficient":
        mapping["values"][0]["coefficient"] = 123  # No silent sign correction.
    elif mutation == "semantic_role":
        _receipt(bundle)["lane_axis"]["lane_keys"][0][1] = "COMPARATIVE_PERIOD"
    else:
        mapping["values"][0]["state"] = "EXACT_COMPONENT_SUM"
    with pytest.raises((ValueError, KeyError)):
        bridge.project_multitable_document_v1(bundle)


def test_blank_stays_none_and_printed_dash_stays_literal_dash() -> None:
    bundle = _bundle()
    _mapping(bundle)["values"][0] = {"coefficient": None, "source_text": None, "state": "BLANK_SOURCE_CELL"}
    _raw(bundle)["values_exact"][0] = None
    payload, _ = _render(bundle)
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    try:
        sheet = workbook["TM"]
        headers = [cell.value for cell in sheet[1]]
        current_col = next(n for n, h in enumerate(headers, 1) if "2025-12-31" in str(h) and str(h).endswith("::Value"))
        prior_col = next(n for n, h in enumerate(headers, 1) if "2024-12-31" in str(h) and str(h).endswith("::Value"))
        assert sheet.cell(3, current_col).value is None
        assert sheet.cell(3, current_col + 1).value == "BLANK_SOURCE_CELL"
        assert sheet.cell(3, prior_col).value == "-"
        assert sheet.cell(3, prior_col).data_type == "s"
    finally:
        workbook.close()


@pytest.mark.parametrize("value", [999999999999999, 1000000000000000, 9007199254740993, -9999999999999999999999999999])
def test_large_signed_integers_have_exact_typed_serialized_readback(value: int) -> None:
    bundle = _bundle()
    raw = str(value)
    _mapping(bundle)["values"][0] = {"coefficient": value, "source_text": raw, "state": "RAW_SIGNED_INTEGER"}
    _raw(bundle)["values_exact"][0] = raw
    payload, _ = _render(bundle)
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    try:
        sheet = workbook["TM"]
        column = next(cell.column for cell in sheet[1] if "2025-12-31" in str(cell.value) and str(cell.value).endswith("::Value"))
        actual = sheet.cell(3, column)
        if len(str(abs(value))) > 15:
            assert actual.value == str(value) and actual.data_type == "s"
        else:
            assert type(actual.value) is int and actual.value == value
    finally:
        workbook.close()


def test_deterministic_literal_text_and_source_only_provenance() -> None:
    first, _ = _render(_bundle())
    second, _ = _render(_bundle())
    assert first == second
    workbook = openpyxl.load_workbook(BytesIO(first), data_only=False)
    try:
        assert workbook["TM"]["B4"].value == "=literal name"
        assert workbook["TM"]["B4"].data_type == "s"
        raw = "".join(str(row[5].value or "") for row in list(workbook["Provenance"].rows)[1:])
        assert "source_only_unmapped_rows" in raw and "retained" in raw
        assert "NOT_CLAIMED" in [cell.value for row in workbook["Status"].rows for cell in row]
    finally:
        workbook.close()


def test_equal_duplicate_slot_is_rejected() -> None:
    bundle = _bundle()
    records = bridge.project_multitable_document_v1(bundle)
    duplicate = copy.deepcopy(records[0])
    duplicate["record_id"] += "other"
    with pytest.raises(ValueError, match="duplicate document/RNID/full-period"):
        renderer.render_projected_tm_workbook_v1(document=bundle["document"], records=[*records, duplicate],
                                                template=[(1249, "cash")], provenance={})


def test_typed_readback_rejects_bool_equals_numeric_zero() -> None:
    workbook = openpyxl.Workbook()
    workbook.active.title = "test"
    workbook.active["A1"] = False
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    with pytest.raises(ValueError, match="typed workbook value drift"):
        renderer.verify_projected_workbook_v1(output.getvalue(), {"test": [[0]]})


@pytest.mark.parametrize("state", ["BLANK_ZERO_AFTER_ONE_UNKNOWN_EQUATION_EXACT", "INFERRED_ONE_UNKNOWN", "BACKSOLVED_SOURCE", "IMPUTED_ZERO"])
def test_legacy_imputation_rejected_even_in_nonselected_candidate(state: str) -> None:
    with pytest.raises(ValueError, match="prohibited"):
        bridge.reject_legacy_imputation_v1({"historical_official": [{"state": state, "coefficient": 0}]})


def test_blank_cannot_be_numeric_zero() -> None:
    with pytest.raises(ValueError, match="blank has a fabricated value"):
        renderer.exact_excel_cell_v1({"coefficient": 0, "source_text": None, "state": "BLANK_SOURCE_CELL"})


def test_report_2024_is_rejected_but_comparative_2024_is_allowed() -> None:
    bundle = _bundle()
    _receipt(bundle)["lane_axis"]["source_lane_keys"] = [["DATE", "2024-12-31"], ["DATE", "2023-12-31"]]
    with pytest.raises(ValueError, match="outside 2025-current"):
        bridge.project_multitable_document_v1(bundle)


def test_nonready_document_is_not_published_as_empty_success() -> None:
    bundle = _bundle()
    bundle["trial"]["status"] = "UNRESOLVED_GEMINI_JSON_FAMILY"
    bundle["trial"]["mappings"] = []
    with pytest.raises(ValueError, match="no accepted-source workbook"):
        bridge.project_multitable_document_v1(bundle)


@pytest.mark.parametrize("payload", [b'{"a":1,"a":2}', b'{"a":NaN}', b'{"a":Infinity}'])
def test_strict_json_rejects_ambiguous_metadata(payload: bytes) -> None:
    with pytest.raises(ValueError):
        bridge.strict_json_v1(payload)


def test_file_pin_and_readonly_sqlite_guard(tmp_path: Path) -> None:
    path = tmp_path / "synthetic.sqlite3"
    with sqlite3.connect(path) as connection:
        connection.execute("CREATE TABLE example(value INTEGER)")
        connection.execute("INSERT INTO example VALUES (1)")
    original = path.read_bytes()
    reference = bridge.file_ref_v1(path)
    connection = bridge._ro(path)
    try:
        assert connection.execute("SELECT value FROM example").fetchone()[0] == 1
        with pytest.raises(sqlite3.OperationalError):
            connection.execute("INSERT INTO example VALUES (2)")
    finally:
        connection.close()
    assert path.read_bytes() == original
    assert bridge.authenticate_ref_v1(reference) == reference
    path.write_bytes(original + b"tamper")
    with pytest.raises(ValueError, match="SHA/byte mismatch"):
        bridge.authenticate_ref_v1(reference)


def test_sqlite_sidecars_are_rejected(tmp_path: Path) -> None:
    path = tmp_path / "example.sqlite3"
    path.write_bytes(b"synthetic")
    Path(str(path) + "-wal").write_bytes(b"")
    with pytest.raises(ValueError, match="sealed snapshot"):
        bridge._ro(path)


def test_actual_callable_origins_cannot_be_replaced_by_head_string(tmp_path: Path) -> None:
    fake = tmp_path / "caller-head.txt"
    fake.write_text("8efd618b6c77f0cdbb402a440e7ba3b3549184f1", encoding="utf-8")
    refs = [bridge.file_ref_v1(fake)]
    with pytest.raises(ValueError, match="loaded runtime producer"):
        bridge._producer_files_v1({"producer_root": str(tmp_path), "producer_refs": refs}, refs)


def test_provenance_chunks_roundtrip_without_excel_truncation() -> None:
    value = {"raw": "=literal" + "x" * 75000}
    rows = renderer._json_rows("TEST", "identity", value)
    assert all(len(row[5]) <= 24000 for row in rows)
    text = "".join(row[5] for row in rows)
    assert json.loads(text) == value
    assert sha256(text.encode("utf-8")).hexdigest() == rows[0][4]


def test_non_bmp_provenance_chunks_fit_excel_utf16_limit() -> None:
    rows = renderer._json_rows("TEST", "non_bmp", {"raw": "\U0001f4b0" * 25000})
    assert all(len(row[5].encode("utf-16-le")) // 2 <= 24000 for row in rows)


def test_every_persisted_trial_mapping_and_candidate_join_is_checked(tmp_path: Path) -> None:
    # Small synthetic store trace projection, not an actual restored database.
    trial = copy.deepcopy(_bundle()["trial"])
    trial["document_ordinal"] = 1
    candidate = trial["candidates"][0]
    candidate.update(page_json_version_id="page:test", physical_page=1,
                     section_id="s1", table_id="t1", status=trial["status"], reasons=[])
    mapping = trial["mappings"][0]
    path = tmp_path / "trace.sqlite3"
    with sqlite3.connect(path) as connection:
        connection.row_factory = sqlite3.Row
        connection.executescript("""
        CREATE TABLE family_trial(family_run_id,document_ordinal,source_logical_name,source_sha256,
          status,candidate_count,selected_candidate_id,mapping_count,reasons_json,trial_sha256,trial_bytes);
        CREATE TABLE family_mapping(family_run_id,document_ordinal,mapping_ordinal,report_norm_id,
          role,row_id,mapping_sha256,mapping_bytes);
        CREATE TABLE family_candidate(family_run_id,document_ordinal,candidate_id,page_json_version_id,
          physical_page,section_id,table_id,status,reason_count,mapping_count,candidate_sha256,candidate_bytes);
        """)
        raw = bridge.json_bytes(trial)
        connection.execute("INSERT INTO family_trial VALUES (?,?,?,?,?,?,?,?,?,?,?)", (
            "run:test", 1, trial["source_logical_name"], trial["source_sha256"], trial["status"], 1,
            trial["selected_candidate_id"], 1, bridge.json_bytes([]), sha256(raw).hexdigest(), raw))
        raw = bridge.json_bytes(mapping)
        connection.execute("INSERT INTO family_mapping VALUES (?,?,?,?,?,?,?,?)", (
            "run:test", 1, 1, mapping["report_norm_id"], mapping["role"], mapping["row_id"], sha256(raw).hexdigest(), raw))
        raw = bridge.json_bytes(candidate)
        connection.execute("INSERT INTO family_candidate VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", (
            "run:test", 1, candidate["candidate_id"], "page:test", 1, "s1", "t1", trial["status"], 0,
            1, sha256(raw).hexdigest(), raw))
        bridge._check_trace_rows_v1(connection, "run:test", {"trials": [trial]})
        connection.execute("UPDATE family_mapping SET report_norm_id=999999")
        with pytest.raises(ValueError, match="mapping projection drift"):
            bridge._check_trace_rows_v1(connection, "run:test", {"trials": [trial]})


def test_publish_creates_nothing_when_authenticated_input_gate_fails(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    output = tmp_path / "must-not-exist"

    def reject(_request: dict) -> None:
        raise ValueError("synthetic upstream authentication failure")

    monkeypatch.setattr(bridge, "authenticate_multitable_store_run_v1", reject)
    with pytest.raises(ValueError, match="upstream authentication failure"):
        bridge.export_tm_multitable_store_document_v1(request={"output_dir": str(output)})
    assert not output.exists()


def test_template_pin_is_fixed_before_workbook_parsing() -> None:
    with pytest.raises(ValueError, match="template SHA mismatch"):
        renderer.template_rows_v1(b"not the pinned workbook")
