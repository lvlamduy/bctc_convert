"""Opt-in strict-core view of the notes pilot workbook (Codex 0357 / 0359). ``export_v1`` and its default workbook
are untouched; this is a separate entry point over the same template, selection, source text, writer and exact readback.

Input contract (the root's pure ledger projection, base 6ce7e1d0):
* every OPERATING_EXPENSE cell carries the additive top-level ``ledger_class`` (one non-empty str) and
  ``ledger_evidence`` (an object); no cell of another family carries either key;
* ``strict_core_eligible`` is the caller's aligned boolean axis (``strict_core_eligibility_v1``): True only for a
  CORE_CANDIDATE whose original acceptance is ACCEPTED (re-checked here; the full eligibility rule is the projection's).
Anything else is refused - there is no silent fallback.

* TM: the pinned template rows and ``export_v1``'s (ReportNormId, period) selection, unchanged. An OPERATING_EXPENSE
  cell shows its number ONLY when strict-core eligible; every other F36 cell keeps its state / acceptance / unit /
  source with no number (printed blanks and dashes stay printed states, never 0). Other families keep the default
  behaviour. One extra ``{period}|ledger_class`` column per period (``NOT_LEDGERED`` outside F36).
* DIMENSIONAL: every cell (observations, backlog, source-only) with the 18 default columns plus ``ledger_class``,
  ``strict_core`` and the canonical ``ledger_evidence`` text for review.
* README: NOTES_STRICT_CORE_VIEW_EXPERIMENTAL; a CORE_CANDIDATE number is not financial acceptance; legacy NOT SUPPLIED;
  legacy / PDF / residual / terminal gates pending.
"""

from __future__ import annotations

import hashlib
import json
from io import BytesIO
from typing import Any

import openpyxl

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes, workbook_has_formula
from bctc_ai.notes.export_v1 import (
    _ABSENT,
    NotesWorkbookError,
    _dimensional_rows_v1,
    _source,
    _tm_cells,
    _verify_sheet,
    _write_sheet,
)
from bctc_ai.primary_statements.excel_v1 import excel_number_v1

STRICT_WORKBOOK_FORMAT = "NOTES_STRICT_CORE_WORKBOOK_V1"
STRICT_PRODUCT_LABEL = "NOTES_STRICT_CORE_VIEW_EXPERIMENTAL"
LEDGERED_FAMILY = "OPERATING_EXPENSE"
NOT_LEDGERED = "NOT_LEDGERED"
_EXCEL_TEXT_LIMIT = 32767


def _check_contract(cells: list[dict[str, Any]], eligible: list[bool]) -> None:
    if type(eligible) is not list or len(eligible) != len(cells) or any(type(flag) is not bool for flag in eligible):
        raise NotesWorkbookError("strict-core eligibility is not one bool per cell")
    for index, (cell, flag) in enumerate(zip(cells, eligible, strict=True)):
        if cell.get("family") == LEDGERED_FAMILY:
            if not (isinstance(cell.get("ledger_class"), str) and cell["ledger_class"]
                    and isinstance(cell.get("ledger_evidence"), dict)):
                raise NotesWorkbookError(f"cell {index}: OPERATING_EXPENSE cell without ledger metadata")
        elif "ledger_class" in cell or "ledger_evidence" in cell:
            raise NotesWorkbookError(f"cell {index}: ledger metadata on a {cell.get('family')} cell")
        if flag and not (cell.get("family") == LEDGERED_FAMILY and cell["ledger_class"] == "CORE_CANDIDATE"
                         and cell.get("acceptance") == "ACCEPTED"):
            raise NotesWorkbookError(f"cell {index}: strict-core eligible but not an ACCEPTED CORE_CANDIDATE")


def _number(cell: dict[str, Any], eligible: bool) -> Any:
    if cell["family"] == LEDGERED_FAMILY:
        return excel_number_v1(cell["value"] if eligible else None)[0]
    return excel_number_v1(cell["value"] if cell["acceptance"] != "HOLD" else None)[0]  # export_v1's default rule


def _evidence_text(cell: dict[str, Any]) -> str | None:
    if "ledger_evidence" not in cell:
        return None
    text = json.dumps(cell["ledger_evidence"], ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
    if len(text) > _EXCEL_TEXT_LIMIT:
        raise NotesWorkbookError("ledger evidence exceeds one Excel text cell; no truncation")
    return text


def _strict_sheets_v1(document, cells, eligible, template, pins) -> tuple[dict[str, list[list[Any]]], int]:
    """The one strict projection both the writer and the verifier use."""

    _check_contract(cells, eligible)
    flags = {id(cell): flag for cell, flag in zip(cells, eligible, strict=True)}
    periods, chosen = _tm_cells(cells)
    header = ["order", "ReportNormId", "name"]
    for period in periods:
        header += [f"{period}|value", f"{period}|state", f"{period}|acceptance", f"{period}|unit", f"{period}|source",
                   f"{period}|ledger_class"]
    tm, numbers = [header], 0
    for order, (rnid, name) in enumerate(template, 1):
        row: list[Any] = [order, rnid, name]
        for period in periods:
            cell = chosen.get((rnid, period))
            if cell is None:
                row += [_ABSENT] * 6
                continue
            number = _number(cell, flags[id(cell)])
            numbers += number is not None
            row += [number, cell["observation_state"], cell["acceptance"], cell["unit"], _source(cell),
                    cell.get("ledger_class", NOT_LEDGERED)]
        tm.append(row)
    dimensional = _dimensional_rows_v1(cells)
    dimensional[0] = dimensional[0] + ["ledger_class", "strict_core", "ledger_evidence"]
    for row, cell in zip(dimensional[1:], cells, strict=True):
        row += [cell.get("ledger_class", NOT_LEDGERED), "ELIGIBLE" if flags[id(cell)] else "NOT_ELIGIBLE",
                _evidence_text(cell)]
    readme = [["product_label", STRICT_PRODUCT_LABEL],
              ["notice", "Experimental strict-core view of the notes pilot. An OPERATING_EXPENSE number appears only for a "
                         "strict-core eligible CORE_CANDIDATE; this is not financial acceptance. Other families keep "
                         "the default pilot behaviour; every other TM row is NOT EVALUATED (not NOT_OBSERVED)."],
              ["legacy", "NOT_SUPPLIED"], ["terminal_financial_ledger", "false"],
              ["pending_gates", "legacy, PDF, residual, terminal"],
              ["document_plan_id", document["document_plan_id"]], ["relative_path", document["relative_path"]],
              ["source_sha256", document["source_sha256"]], ["families", ",".join(sorted({c["family"] for c in cells}))],
              ["pins", repr(sorted(pins.items()))]]
    return {"README": readme, "TM": tm, "DIMENSIONAL": dimensional}, numbers


def build_notes_strict_core_workbook_v1(document: dict[str, Any], cells: list[dict[str, Any]], strict_core_eligible: list[bool],
                                        template: list[tuple[int, str]], pins: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
    workbook = openpyxl.Workbook()
    sheets, _numbers = _strict_sheets_v1(document, cells, strict_core_eligible, template, pins)
    workbook.active.title = "README"
    _write_sheet(workbook.active, sheets["README"])
    for name in ("TM", "DIMENSIONAL"):
        _write_sheet(workbook.create_sheet(name), sheets[name])
    payload = deterministic_workbook_bytes(workbook, creator="bctc-ai notes strict-core view v1")
    receipt = verify_notes_strict_core_workbook_v1(payload, document, cells, strict_core_eligible, template, pins)
    return payload, receipt


def verify_notes_strict_core_workbook_v1(payload: bytes, document: dict[str, Any], cells: list[dict[str, Any]],
                                         strict_core_eligible: list[bool], template: list[tuple[int, str]],
                                         pins: dict[str, Any]) -> dict[str, Any]:
    """Every emitted cell reads back with its exact value AND type, no extra cell, no formula; a number smuggled into a
    non-eligible F36 position is a readback difference."""

    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    if workbook_has_formula(workbook):
        raise NotesWorkbookError("workbook contains a formula")
    if workbook.sheetnames != ["README", "TM", "DIMENSIONAL"]:
        raise NotesWorkbookError("workbook sheet set drifted")
    sheets, numbers = _strict_sheets_v1(document, cells, strict_core_eligible, template, pins)
    compared = sum(_verify_sheet(workbook[name], rows, name) for name, rows in sheets.items())
    return {"format_version": STRICT_WORKBOOK_FORMAT, "workbook_sha256": hashlib.sha256(payload).hexdigest(),
            "workbook_size_bytes": len(payload), "tm_numbers_emitted": numbers, "dimensional_rows": len(cells),
            "strict_core_eligible": sum(strict_core_eligible), "grid_positions_compared": compared,
            "readback_scope": "EVERY_EMITTED_CELL_TYPED_NO_EXTRA_CELLS", "formula_cells": 0,
            "financial_acceptance": False, "status": "PASS_EXACT_READBACK"}
