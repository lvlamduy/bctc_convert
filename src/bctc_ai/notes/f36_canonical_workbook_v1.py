"""TM-template workbook of one document's persisted F36 canonical proposal (EXPERIMENTAL; never financial acceptance).

Numbers come ONLY from the canonical mappings of the document's trial (the B adapter's admitted, source-replayed
pairs): one TM value per (ReportNormId, period) in the mapping's own CURRENT / COMPARATIVE order, as the exact signed
coefficient in the declared table unit, with its ORIGINAL source reference. Every notes OPEX cell stays on DIMENSIONAL
with its ledger class and canonical status (mapped, or the adapter's named withhold reason): nothing is dropped,
nothing is zero-filled, no amount is derived. The caller authenticates the store / sweep it loaded the trial from.
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from typing import Any

import openpyxl

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes, workbook_has_formula
from bctc_ai.notes import export_v1 as pilot
from bctc_ai.source_structure.contracts_v1 import canonical_json_bytes_v1

FORMAT = "F36_CANONICAL_PROPOSAL_WORKBOOK_V1"
PRODUCT_LABEL = "EXPERIMENTAL_CANONICAL_PROPOSAL_NOT_ACCEPTANCE"
_TM_FIELDS = ("value", "unit", "role", "source")


class F36CanonicalWorkbookError(pilot.NotesWorkbookError):
    """The canonical proposal workbook cannot be projected or read back exactly."""


def _need(condition: bool, message: str) -> None:
    if not condition:
        raise F36CanonicalWorkbookError(message)


def _source(ref: dict[str, Any]) -> str:
    loc = ref["locator"]
    return (f"p{loc['physical_page']}/{loc['section_id']}/{loc['table_id']}/{ref['row_id']}"
            f"/col{ref['money_column_ordinals'][0]}@{loc['page_json_version_id']}")


def _key(cell: dict[str, Any]) -> bytes:
    return canonical_json_bytes_v1(cell)


def _tm_slots(trial: dict[str, Any], template_ids: set[int]) -> dict[tuple[int, str], tuple[Any, ...]]:
    slots = {}
    for mapping in trial["mappings"]:
        rnid = mapping["report_norm_id"]
        _need(type(rnid) is int and rnid in template_ids, f"Canonical RNID {rnid} is outside the supplied template")
        _need(len(mapping["values"]) == len(mapping["period_axis"]) == len(mapping["source_refs"]) == 2,
              "Canonical mapping value/period/source axes differ")
        for lane, (value, period, ref) in enumerate(zip(mapping["values"], mapping["period_axis"], mapping["source_refs"],
                                                         strict=True)):
            _need(type(value.get("coefficient")) is int and value.get("state") == "RAW_SIGNED_INTEGER"
                  and period["period_key"] == ref["period_key"], "Canonical lane is not an exact printed integer")
            slot = (rnid, period["period_key"])
            _need(slot not in slots, f"Two canonical values for ReportNormId {rnid} period {period['period_key']}")
            role = period["parsed_period"]["role"]
            _need(role == ("CURRENT" if lane == 0 else "COMPARATIVE"), "Canonical lane order is not CURRENT then COMPARATIVE")
            slots[slot] = (pilot.excel_number_v1(str(value["coefficient"]))[0], mapping["unit"], role, _source(ref))
    return slots


def _statuses(trial: dict[str, Any]) -> dict[bytes, str]:
    ledger, statuses = trial["notes_ledger"], {}
    mapped = {(ref["locator"]["page_json_version_id"], ref["locator"]["section_id"], ref["locator"]["table_id"],
               ref["row_ordinal"], ref["money_column_ordinals"][0])
              for mapping in trial["mappings"] for ref in mapping["source_refs"]}
    for record in ledger["retained"]:
        cell = record["cell"] if "cell" in record else record["classification"]["cell"]
        statuses[_key(cell)] = "RETAINED:" + record["adapter_disposition"]
    for cell in ledger["cells"]:
        origin = cell.get("original_cell") or {}
        coordinate = (origin.get("page_json_version_id"), f"s{origin.get('section_ordinal')}", f"t{origin.get('table_ordinal')}",
                      origin.get("row_ordinal"), origin.get("column_ordinal"))
        if coordinate in mapped:
            _need(_key(cell) not in statuses, "A canonical mapped cell is also retained")
            statuses[_key(cell)] = "CANONICAL_MAPPED"
    return statuses


def _sheets_v1(document, trial, template, pins):
    _need(type(document.get("document_plan_id")) is str and bool(document["document_plan_id"]), "Document identity")
    cells, records = trial["notes_ledger"]["cells"], trial["notes_ledger"]["classifications"]
    _need(len(cells) == len(records) and all(_key(r["cell"]) == _key(c) for c, r in zip(cells, records, strict=True)),
          "Ledger cell/classification axis differs")
    _need(all(c.get("document_plan_id") == document["document_plan_id"] for c in cells), "Cells belong to another document")
    template_ids = {rnid for rnid, _name in template}
    _need(len(template_ids) == len(template), "Template repeats a ReportNormId")
    slots = _tm_slots(trial, template_ids)
    keys = {period for _rnid, period in slots}
    dated = sorted((p for p in keys if pilot._period_order(p)[0] == 0), key=lambda p: (pilot._period_order(p)[1], p), reverse=True)
    periods = dated + sorted(p for p in keys if pilot._period_order(p)[0] == 1)
    header = ["order", "ReportNormId", "name"] + [f"{period}|{field}" for period in periods for field in _TM_FIELDS]
    tm, placed = [header], 0
    for order, (rnid, name) in enumerate(template, 1):
        row: list[Any] = [order, rnid, name]
        for period in periods:
            slot = slots.get((rnid, period))
            row += list(slot) if slot else [pilot._ABSENT] * len(_TM_FIELDS)
            placed += slot is not None
        tm.append(row)
    statuses = _statuses(trial)
    dimensional = pilot._dimensional_rows_v1(cells)
    dimensional[0] = dimensional[0] + ["ledger_class", "canonical_status"]
    for line, cell, record in zip(dimensional[1:], cells, records, strict=True):
        line += [record["class"], statuses.get(_key(cell), "NOT_CANONICAL_NO_RETAINED_RECORD")]
    readme = [["product_label", PRODUCT_LABEL],
              ["notice", "F36 TM numbers are the persisted canonical B proposal (source-replayed, ROOT-closed, exact printed "
                         "integers in the declared table unit). Not financial acceptance or a family seal. Every other notes "
                         "observation stays on DIMENSIONAL with its ledger class and withhold reason."],
              ["legacy_comparison", "NOT_SUPPLIED"], ["input_authentication", "CALLER_BOUNDARY"],
              ["document_plan_id", document["document_plan_id"]], ["relative_path", document["relative_path"]],
              ["source_sha256", document["source_sha256"]], ["canonical_mappings", len(trial["mappings"])],
              ["pins", repr(sorted(pins.items()))]]
    return {"README": readme, "TM": tm, "DIMENSIONAL": dimensional}, placed


def build_f36_canonical_workbook_v1(document: dict[str, Any], trial: dict[str, Any], template: list[tuple[int, str]],
                                    pins: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
    sheets, _placed = _sheets_v1(document, trial, template, pins)
    workbook = openpyxl.Workbook()
    workbook.active.title = "README"
    pilot._write_sheet(workbook.active, sheets["README"])
    for name in ("TM", "DIMENSIONAL"):
        pilot._write_sheet(workbook.create_sheet(name), sheets[name])
    payload = deterministic_workbook_bytes(workbook, creator="bctc-ai F36 canonical proposal v1")
    return payload, verify_f36_canonical_workbook_v1(payload, document, trial, template, pins)


def verify_f36_canonical_workbook_v1(payload: bytes, document: dict[str, Any], trial: dict[str, Any],
                                     template: list[tuple[int, str]], pins: dict[str, Any]) -> dict[str, Any]:
    sheets, placed = _sheets_v1(document, trial, template, pins)
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    _need(not workbook_has_formula(workbook), "Canonical workbook contains a formula")
    _need(workbook.sheetnames == ["README", "TM", "DIMENSIONAL"], "Canonical workbook sheet set differs")
    compared = sum(pilot._verify_sheet(workbook[name], rows, name) for name, rows in sheets.items())
    return {"format_version": FORMAT, "product_label": PRODUCT_LABEL, "status": "PASS_EXACT_READBACK",
            "workbook_sha256": hashlib.sha256(payload).hexdigest(), "workbook_size_bytes": len(payload),
            "tm_values": placed, "canonical_mappings": len(trial["mappings"]),
            "dimensional_rows": len(trial["notes_ledger"]["cells"]), "grid_positions_compared": compared,
            "readback_scope": "EVERY_EMITTED_CELL_TYPED_NO_EXTRA_CELLS", "formula_cells": 0,
            "financial_acceptance": False}
