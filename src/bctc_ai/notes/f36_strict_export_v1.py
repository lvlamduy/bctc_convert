"""Opt-in experimental F36 ledger-gated Excel; existing exporter is unchanged.

The caller authenticates template, store, ledger and source artifacts. This byte
builder/verifier authenticates none of those files. TM numbers are candidates,
never financial acceptance; DIMENSIONAL retains all original observations.
"""

from __future__ import annotations

import hashlib
from collections import defaultdict
from io import BytesIO
from typing import Any

import openpyxl

from bctc_ai.export.canonical_xlsx import deterministic_workbook_bytes, workbook_has_formula
from bctc_ai.notes import export_v1 as pilot
from bctc_ai.notes.f36_ledger_projection_v1 import FAMILY, strict_f36_core_eligibility_v1
from bctc_ai.source_structure.contracts_v1 import canonical_json_bytes_v1

FORMAT = "F36_STRICT_EXPERIMENTAL_NOTES_WORKBOOK_V1"
PRODUCT_LABEL = "EXPERIMENTAL_CORE_CANDIDATE_NOT_ACCEPTANCE"


class StrictF36WorkbookError(pilot.NotesWorkbookError):
    """A strict candidate workbook cannot be projected or read back exactly."""


def _need(condition: bool, message: str) -> None:
    if not condition:
        raise StrictF36WorkbookError(message)


def _json(value: Any) -> str:
    return canonical_json_bytes_v1(value).decode("utf-8").removesuffix("\n")


def _tm_v1(cells, template, eligibility):
    # Other families retain the existing default selection and number policy.
    _periods, foreign = pilot._tm_cells([c for c in cells if c["family"] != FAMILY])
    selected = {key: (cell, None) for key, cell in foreign.items()}
    groups, seen = defaultdict(list), set()
    template_ids = {rnid for rnid, _name in template}
    _need(len(template_ids) == len(template), "Template repeats an RNID")
    for index, (cell, eligible) in enumerate(zip(cells, eligibility, strict=True)):
        if cell["family"] != FAMILY:
            continue
        encoded = canonical_json_bytes_v1(cell)
        _need(encoded not in seen, "Duplicate exact F36 observation")
        seen.add(encoded)
        rnid, period = cell.get("report_norm_id"), cell.get("period_key")
        _need(not eligible or (type(rnid) is int and type(period) is str and bool(period)),
              "Core candidate has no complete template period key")
        if rnid is None or not period:
            continue  # Still present in DIMENSIONAL, including source-only members.
        _need(type(rnid) is int and rnid in template_ids and type(period) is str,
              "F36 mapped observation is outside the supplied template/period axis")
        groups[(rnid, period)].append((index, cell, eligible))
    numeric = 0
    for key, group in groups.items():
        _need(key not in selected, "F36 and another family share an RNID/period slot")
        core = [item for item in group if item[2]]
        _need(len(core) <= 1, "Multiple core candidates for one RNID/period")
        if core:
            selected[key] = (core[0][1], True)
            numeric += 1
        elif len(group) == 1:
            selected[key] = (group[0][1], False)
        else:
            # No selected source is invented. Each original state/source remains
            # literal here and every complete observation remains in DIMENSIONAL.
            selected[key] = (group, "MULTIPLE_NONCORE")
    keys = {period for _rnid, period in selected}
    dated = sorted((p for p in keys if pilot._period_order(p)[0] == 0),
                   key=lambda p: (pilot._period_order(p)[1], p), reverse=True)
    periods = dated + sorted(p for p in keys if pilot._period_order(p)[0] == 1)
    header = ["order", "ReportNormId", "name"]
    for period in periods:
        header += [f"{period}|{field}" for field in ("value", "state", "acceptance", "unit", "source", "ledger_class")]
    rows, placed = [header], 0
    for order, (rnid, name) in enumerate(template, 1):
        row = [order, rnid, name]
        for period in periods:
            chosen = selected.get((rnid, period))
            if chosen is None:
                row += [pilot._ABSENT] * 6
                continue
            item, gate = chosen
            if gate == "MULTIPLE_NONCORE":
                row += [None, _json({"multiplicity": len(item), "states": [c["observation_state"] for _i, c, _e in item]}),
                        _json([c["acceptance"] for _i, c, _e in item]), _json([c["unit"] for _i, c, _e in item]),
                        _json([{"dimensional_row": i + 2, "source": pilot._source(c)} for i, c, _e in item]),
                        _json([c["ledger_class"] for _i, c, _e in item])]
            else:
                value = item["value"] if (gate is True or (gate is None and item["acceptance"] != "HOLD")) else None
                number, _policy = pilot.excel_number_v1(value)
                row += [number, item["observation_state"], item["acceptance"], item["unit"], pilot._source(item),
                        item.get("ledger_class")]
            placed += 1
        rows.append(row)
    return rows, placed, numeric


def _sheets_v1(document, cells, template, pins):
    eligibility = strict_f36_core_eligibility_v1(cells)  # Validates ALL F36 attachments before selecting anything.
    _need(type(document.get("document_plan_id")) is str and bool(document["document_plan_id"])
          and all(c.get("document_plan_id") == document["document_plan_id"] for c in cells),
          "Workbook cells do not name this one document plan")
    tm, placed, numeric = _tm_v1(cells, template, eligibility)
    readme = pilot._readme_rows_v1(document, cells, pins)
    readme[0] = ["product_label", PRODUCT_LABEL]
    readme[1] = ["notice", "F36 TM numbers are ledger CORE_CANDIDATE only, not canonical/financial acceptance. "
               "Other families retain the existing experimental exporter policy. All observations remain DIMENSIONAL."]
    readme += [["legacy_comparison", "NOT_SUPPLIED_BY_THIS_WRAPPER"],
               ["input_authentication", "CALLER_BOUNDARY_NOT_PERFORMED_BY_WRAPPER"],
               ["family_seal_or_canonical_release", False]]
    dimensional = pilot._dimensional_rows_v1(cells)
    dimensional[0] += ["ledger_class", "ledger_evidence"]
    for line, cell in zip(dimensional[1:], cells, strict=True):
        line += [cell.get("ledger_class"), _json(cell["ledger_evidence"]) if "ledger_evidence" in cell else None]
    return {"README": readme, "TM": tm, "DIMENSIONAL": dimensional}, placed, numeric


def build_strict_f36_workbook_v1(document: dict[str, Any], cells: list[dict[str, Any]],
                                 template: list[tuple[int, str]], pins: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
    """Use the caller's pinned tm_template_rows_v1 result; return bytes, never write files or mutate input cells."""

    sheets, _placed, _numeric = _sheets_v1(document, cells, template, pins)
    workbook = openpyxl.Workbook()
    workbook.active.title = "README"
    pilot._write_sheet(workbook.active, sheets["README"])
    for name in ("TM", "DIMENSIONAL"):
        pilot._write_sheet(workbook.create_sheet(name), sheets[name])
    payload = deterministic_workbook_bytes(workbook, creator="bctc-ai strict F36 candidate v1")
    return payload, verify_strict_f36_workbook_v1(payload, document, cells, template, pins)


def verify_strict_f36_workbook_v1(payload: bytes, document: dict[str, Any], cells: list[dict[str, Any]],
                                  template: list[tuple[int, str]], pins: dict[str, Any]) -> dict[str, Any]:
    """Typed readback of every emitted cell, including literal evidence and all metadata; not source authentication."""

    sheets, placed, numeric = _sheets_v1(document, cells, template, pins)
    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    _need(not workbook_has_formula(workbook), "Strict workbook contains a formula")
    _need(workbook.sheetnames == ["README", "TM", "DIMENSIONAL"], "Strict workbook sheet set differs")
    compared = sum(pilot._verify_sheet(workbook[name], rows, name) for name, rows in sheets.items())
    return {"format_version": FORMAT, "product_label": PRODUCT_LABEL, "status": "PASS_EXACT_READBACK",
            "workbook_sha256": hashlib.sha256(payload).hexdigest(), "workbook_size_bytes": len(payload),
            "tm_cells_checked": placed, "strict_f36_numeric_cells": numeric, "dimensional_rows": len(cells),
            "grid_positions_compared": compared, "readback_scope": "EVERY_EMITTED_CELL_TYPED_NO_EXTRA_CELLS",
            "formula_cells": 0, "source_or_ledger_file_authentication_performed": False,
            "canonical_acceptance_or_family_seal_granted": False}
