"""Notes pilot workbook: legacy TM template rows plus the dimensional view, exact readback.

* TM: the rows of the pinned ``Bank_TM_ReportNormId.v2.xlsx`` in template order;
  for every period, value / observation state / acceptance / unit / source
  reference. Only cells mapped through the crosswalk appear here; a held cell
  shows its state and acceptance with no value.
* DIMENSIONAL: every evaluated cell by (concept, axis coordinate, basis,
  period) — the cross-bank comparable view, including cells without a legacy
  ReportNormId.
* README: product label NOTES_PILOT_EXPERIMENTAL, families included, pins.

Numbers follow the PRIMARY exact-number policy; text is always literal, so no
cell can become a formula. ``verify_notes_workbook_v1`` reopens the bytes and
compares every value, type, state, unit, period and source reference.
"""

from __future__ import annotations

import hashlib
import re
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl

from bctc_ai.export.canonical_xlsx import (
    deterministic_workbook_bytes,
    set_literal_cell,
    workbook_has_formula,
)
from bctc_ai.primary_statements.excel_v1 import excel_number_v1
from bctc_ai.primary_statements.source_v1 import sha256_file
from bctc_ai.schema.xlsx_reader import read_rows

NOTES_WORKBOOK_FORMAT = "NOTES_PILOT_WORKBOOK_V2"  # V2: DIMENSIONAL semantic_state column
PRODUCT_LABEL = "NOTES_PILOT_EXPERIMENTAL"
TM_TEMPLATE = "template/Bank_TM_ReportNormId.v2.xlsx"
TM_TEMPLATE_SHA256 = "64589f7ee2b025ac0bd784c13e0b0ba0aff9d46456a83fa7fe2d68a7de375115"
EXPORTED = {"ACCEPTED", "ACCEPTED_DASH", "ACCEPTED_BLANK", "HOLD"}
_DIMENSIONAL_HEADER = ["family", "concept", "axis", "member", "classification_basis", "period_key", "period_basis",
                       "unit", "unit_basis", "observation_state", "value", "raw", "acceptance", "hold_reasons",
                       "report_norm_id", "label_exact", "source", "semantic_state"]


class NotesWorkbookError(ValueError):
    """The notes workbook does not read back exactly."""


def tm_template_rows_v1(project_root: Path) -> list[tuple[int, str]]:
    path = project_root / TM_TEMPLATE
    if sha256_file(path) != TM_TEMPLATE_SHA256:
        raise NotesWorkbookError("TM template SHA-256 differs from the pinned template")
    rows = []
    for row in read_rows(path):
        raw = (row.get("B") or "").strip()
        if raw and raw != "ReportNormId":
            rows.append((int(raw), row.get("C") or ""))
    if len({rnid for rnid, _ in rows}) != len(rows):
        raise NotesWorkbookError("TM template repeats a ReportNormId")
    return rows


def _period_order(key: str) -> tuple[int, str]:
    dates = re.findall(r"\d{4}-\d{2}-\d{2}", key)
    return (0, "".join(dates)) if dates else (1, key)


def _source(cell: dict[str, Any]) -> str:
    loc = cell["locator"]
    return f"p{loc['physical_page']}/{loc['section_id']}/{loc['table_id']}/{loc['row_id']}/{loc['column_id']}@{loc['page_json_version_id']}"


def _tm_cells(cells: list[dict[str, Any]]) -> tuple[list[str], dict[tuple[int, str], dict[str, Any]]]:
    chosen: dict[tuple[int, str], dict[str, Any]] = {}
    for cell in cells:
        if cell.get("report_norm_id") is None or not cell["period_key"] or cell["acceptance"] not in EXPORTED:
            continue
        key = (cell["report_norm_id"], cell["period_key"])
        current = chosen.get(key)
        if current is None:
            chosen[key] = cell
        elif current["acceptance"] != "HOLD" and cell["acceptance"] != "HOLD":
            raise NotesWorkbookError(f"two exported cells for ReportNormId {key[0]} period {key[1]}")
        elif current["acceptance"] == "HOLD" and cell["acceptance"] != "HOLD":
            chosen[key] = cell
    keys = {period for _, period in chosen}
    dated = sorted((p for p in keys if _period_order(p)[0] == 0), key=lambda p: _period_order(p)[1], reverse=True)
    return dated + sorted(p for p in keys if _period_order(p)[0] == 1), chosen


_ABSENT = object()  # a grid position the writer never creates (no TM cell for that ReportNormId and period)


def _readme_rows_v1(document: dict[str, Any], cells: list[dict[str, Any]], pins: dict[str, Any]) -> list[list[Any]]:
    return [["product_label", PRODUCT_LABEL],
            ["notice", "Experimental notes pilot (thuyết minh). Not a production or accepted export; only the "
                       "families listed are evaluated, every other TM row is NOT EVALUATED (not NOT_OBSERVED)."],
            ["document_plan_id", document["document_plan_id"]], ["relative_path", document["relative_path"]],
            ["source_sha256", document["source_sha256"]], ["families", ",".join(sorted({c["family"] for c in cells}))],
            ["pins", repr(sorted(pins.items()))]]


def _tm_rows_v1(cells: list[dict[str, Any]], template: list[tuple[int, str]]) -> tuple[list[list[Any]], int]:
    periods, chosen = _tm_cells(cells)
    header = ["order", "ReportNormId", "name"]
    for period in periods:
        header += [f"{period}|value", f"{period}|state", f"{period}|acceptance", f"{period}|unit", f"{period}|source"]
    rows, placed = [header], 0
    for order, (rnid, name) in enumerate(template, 1):
        row: list[Any] = [order, rnid, name]
        for period in periods:
            cell = chosen.get((rnid, period))
            if cell is None:
                row += [_ABSENT] * 5
                continue
            number, _policy = excel_number_v1(cell["value"] if cell["acceptance"] != "HOLD" else None)
            row += [number, cell["observation_state"], cell["acceptance"], cell["unit"], _source(cell)]
            placed += 1
        rows.append(row)
    return rows, placed


def _dimensional_rows_v1(cells: list[dict[str, Any]]) -> list[list[Any]]:
    return [list(_DIMENSIONAL_HEADER)] + [
        [cell["family"], cell["concept"], cell["axis"], cell["member"], cell["classification_basis"],
         cell["period_key"], cell["period_basis"], cell["unit"], cell["unit_basis"], cell["observation_state"],
         cell["value"], cell["raw"], cell["acceptance"], ";".join(cell["hold_reasons"]),
         cell.get("report_norm_id"), cell["label_exact"], _source(cell),
         # e.g. BASE_OR_TAX_EFFECT_UNPROVEN: an observation, not a proved legacy amount (Codex 0133).
         cell.get("semantic_state")]
        for cell in cells]


def _sheets_v1(document, cells, template, pins) -> tuple[dict[str, list[list[Any]]], int]:
    """The one projection both the writer and the verifier use: every emitted cell of every sheet (Codex 0222)."""

    tm, placed = _tm_rows_v1(cells, template)
    return {"README": _readme_rows_v1(document, cells, pins), "TM": tm, "DIMENSIONAL": _dimensional_rows_v1(cells)}, placed


def _write_sheet(sheet, rows: list[list[Any]]) -> None:
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            if value is _ABSENT:
                continue
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                sheet.cell(r, c).value = value  # exact-number policy (excel_number_v1); ordinals and ReportNormIds
            else:
                set_literal_cell(sheet.cell(r, c), value)  # text is always literal: never a formula


def _as_read(value: Any) -> Any:
    """What a written value reads back as: openpyxl writes an empty string as a value-less cell (None)."""

    return None if value is _ABSENT or value == "" else value


def _verify_sheet(sheet, rows: list[list[Any]], name: str) -> int:
    for r, row in enumerate(rows, 1):
        for c, value in enumerate(row, 1):
            got, expected = sheet.cell(r, c).value, _as_read(value)
            if type(got) is not type(expected) or got != expected:
                raise NotesWorkbookError(f"{name}!R{r}C{c} does not read back exactly")
    for line in sheet.iter_rows():
        for cell in line:
            if cell.value is not None and (cell.row > len(rows) or cell.column > len(rows[cell.row - 1])):
                raise NotesWorkbookError(f"{name}!R{cell.row}C{cell.column} is an extra cell")
    return sum(len(row) for row in rows)


def build_notes_workbook_v1(document: dict[str, Any], cells: list[dict[str, Any]], template: list[tuple[int, str]],
                            pins: dict[str, Any]) -> tuple[bytes, dict[str, Any]]:
    workbook = openpyxl.Workbook()
    sheets, _placed = _sheets_v1(document, cells, template, pins)
    workbook.active.title = "README"
    _write_sheet(workbook.active, sheets["README"])
    for name in ("TM", "DIMENSIONAL"):
        _write_sheet(workbook.create_sheet(name), sheets[name])
    payload = deterministic_workbook_bytes(workbook, creator="bctc-ai notes pilot v1")
    receipt = verify_notes_workbook_v1(payload, document, cells, template, pins)
    return payload, receipt


def verify_notes_workbook_v1(payload: bytes, document: dict[str, Any], cells: list[dict[str, Any]],
                             template: list[tuple[int, str]], pins: dict[str, Any]) -> dict[str, Any]:
    """Every emitted cell of README, TM (header, order, name, every period column) and DIMENSIONAL (all 18 columns)
    reads back with its exact value AND type; any value outside that grid is an extra cell (Codex 0222)."""

    workbook = openpyxl.load_workbook(BytesIO(payload), data_only=False)
    if workbook_has_formula(workbook):
        raise NotesWorkbookError("workbook contains a formula")
    if workbook.sheetnames != ["README", "TM", "DIMENSIONAL"]:
        raise NotesWorkbookError("workbook sheet set drifted")
    sheets, placed = _sheets_v1(document, cells, template, pins)
    compared = sum(_verify_sheet(workbook[name], rows, name) for name, rows in sheets.items())
    return {"format_version": NOTES_WORKBOOK_FORMAT, "workbook_sha256": hashlib.sha256(payload).hexdigest(),
            "workbook_size_bytes": len(payload), "tm_cells_checked": placed, "dimensional_rows": len(cells),
            "grid_positions_compared": compared, "readback_scope": "EVERY_EMITTED_CELL_TYPED_NO_EXTRA_CELLS",
            "formula_cells": 0, "status": "PASS_EXACT_READBACK"}
