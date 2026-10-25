"""TM-template workbook of one document's Family-37 READY proposal (EXPERIMENTAL; never financial acceptance).

Numbers come ONLY from the trial's mappings (the reviewed Family-37 adapter, source replayed by the caller). Family-37
mappings carry two semantic lanes, CURRENT_PERIOD and COMPARATIVE_PERIOD, not dated period keys, so the TM sheet has
one column group per lane and the document's own report period stays in RUN_METADATA. Per lane:

* a printed integer (RAW_SIGNED_INTEGER) is written as that exact signed coefficient;
* a printed dash (DASH_ZERO) is written as the literal "-" - never as 0;
* an exact derivation (component / frontier / same-role sums) is written with its derivation state and the source
  rows it was derived from;
* an unobserved lane or a blank source cell is left EMPTY with its state - nothing is zero-filled or inferred.

The caller authenticates the store / sweep it loaded the trial from.
"""

from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from typing import Any

import openpyxl

from bctc_ai.export.canonical_xlsx import (
    append_literal_row,
    deterministic_workbook_bytes,
    workbook_has_formula,
)
from bctc_ai.source_structure.contracts_v1 import canonical_json_bytes_v1

FORMAT = "F37_CANONICAL_PROPOSAL_WORKBOOK_V1"
PRODUCT_LABEL = "EXPERIMENTAL_CANONICAL_PROPOSAL_NOT_ACCEPTANCE"
TM_TEMPLATE = "template/Bank_TM_ReportNormId.v2.xlsx"
TM_TEMPLATE_SHA256 = "64589f7ee2b025ac0bd784c13e0b0ba0aff9d46456a83fa7fe2d68a7de375115"
READY = "READY_FOR_SCHEMA_MAPPING_REVIEW_PROPOSAL_ONLY"
LANES = ("CURRENT_PERIOD", "COMPARATIVE_PERIOD")
_FIELDS = ("value", "state", "unit", "role", "source")
_PRINTED = "RAW_SIGNED_INTEGER"
_DASH = "DASH_ZERO"
_EMPTY = {"UNOBSERVED_SOURCE_LANE", "BLANK_SOURCE_CELL"}
_DERIVED = {"EXACT_DECLARED_SOURCE_COMPONENT_SUM", "EXACT_COMPLETE_VISIBLE_SCOPED_CHILD_FRONTIER_SUM",
            "EXACT_VISIBLE_CUSTOMER_COMPONENT_SUM", "EXACT_VISIBLE_GENERAL_AND_SPECIFIC_SUM",
            "EXACT_SAME_ROLE_SOURCE_ROW_SUM"}


class F37CanonicalWorkbookError(ValueError):
    """The Family-37 proposal workbook cannot be projected or read back exactly."""


def _need(condition: bool, message: str) -> None:
    if not condition:
        raise F37CanonicalWorkbookError(message)


def tm_template_rows_v1(project_root: Path) -> list[tuple[int, str]]:
    """(ReportNormId, name) rows of the pinned TM template, in template order."""

    path = project_root / TM_TEMPLATE
    _need(hashlib.sha256(path.read_bytes()).hexdigest() == TM_TEMPLATE_SHA256, "TM template SHA-256 differs")
    workbook = openpyxl.load_workbook(path, read_only=True)
    try:
        rows = [(int(row[1]), str(row[2] or "")) for row in workbook.worksheets[0].iter_rows(min_row=2, values_only=True)
                if row and row[1] is not None]
    finally:
        workbook.close()
    _need(len({rnid for rnid, _ in rows}) == len(rows), "TM template repeats a ReportNormId")
    return rows


def _source(ref: dict[str, Any]) -> str:
    loc = ref["locator"]
    columns = ",".join(str(ordinal) for ordinal in ref.get("money_column_ordinals") or [])
    return f"p{loc['physical_page']}/{loc['section_id']}/{loc['table_id']}/r{ref['row_ordinal']}/col{columns}@{loc['page_json_version_id']}"


def _cell(value: dict[str, Any]) -> tuple[Any, str]:
    state, coefficient = value.get("state"), value.get("coefficient")
    if state == _PRINTED:
        _need(type(coefficient) is int and value.get("source_text") is not None, "Printed lane lacks its exact integer")
        return coefficient, state
    if state == _DASH:
        _need(coefficient == 0 and value.get("source_text") is not None and value["source_text"].strip() in {"-", "_", "–"},
              "Dash lane is not a printed dash")
        return "-", state
    if state in _EMPTY:
        _need(coefficient is None, "An empty lane carries a number")
        return None, state
    _need(state in _DERIVED and type(coefficient) is int, f"Unsupported lane state {state}")
    return coefficient, state


def _tm_slots(trial: dict[str, Any], template_ids: set[int]) -> dict[tuple[int, str], tuple[Any, ...]]:
    _need(trial["status"] == READY and trial["mappings"], "Only a READY trial with mappings is exported")
    slots = {}
    for mapping in trial["mappings"]:
        rnid = mapping["report_norm_id"]
        _need(type(rnid) is int and rnid in template_ids, f"ReportNormId {rnid} is outside the supplied template")
        _need(len(mapping["values"]) == len(LANES) and mapping["source_refs"], "Mapping lane / source axes are invalid")
        sources = ";".join(_source(ref) for ref in mapping["source_refs"])
        for lane, value in zip(LANES, mapping["values"], strict=True):
            _need((rnid, lane) not in slots, f"Two values for ReportNormId {rnid} {lane}")
            number, state = _cell(value)
            slots[(rnid, lane)] = (number, state, mapping["unit"], mapping["role"], sources)
    return slots


def _sheets_v1(document: dict[str, Any], trial: dict[str, Any], template: list[tuple[int, str]],
               head: str) -> dict[str, list[list[Any]]]:
    slots = _tm_slots(trial, {rnid for rnid, _ in template})
    tm = [["order", "ReportNormId", "name"] + [f"{lane}|{field}" for lane in LANES for field in _FIELDS]]
    for order, (rnid, name) in enumerate(template, 1):
        row = [order, rnid, name]
        for lane in LANES:
            slot = slots.get((rnid, lane))
            row += list(slot) if slot else [None] * len(_FIELDS)
        tm.append(row)
    source = [["ReportNormId", "role", "mapping_state", "source", "label_exact", "row_kind"]]
    for mapping in sorted(trial["mappings"], key=lambda m: (m["report_norm_id"], m["role"])):
        for ref in mapping["source_refs"]:
            source.append([mapping["report_norm_id"], mapping["role"], mapping["state"], _source(ref),
                           ref.get("label_exact"), ref.get("row_kind")])
    metadata = [["key", "value"], ["format", FORMAT], ["product_label", PRODUCT_LABEL], ["producer_commit", head],
                ["document_ordinal", trial["document_ordinal"]], ["relative_path", document["relative_path"]],
                ["source_sha256", document["source_sha256"]], ["lanes", "CURRENT_PERIOD|COMPARATIVE_PERIOD of the report"],
                ["mappings", len(trial["mappings"])], ["template_sha256", TM_TEMPLATE_SHA256],
                ["trial_sha256", hashlib.sha256(canonical_json_bytes_v1(trial)).hexdigest()]]
    return {"TM": tm, "SOURCE": source, "RUN_METADATA": metadata}


def build_f37_canonical_workbook_v1(document: dict[str, Any], trial: dict[str, Any], template: list[tuple[int, str]],
                                    *, head: str) -> tuple[bytes, dict[str, Any]]:
    """(xlsx bytes, receipt) for one READY trial."""

    sheets = _sheets_v1(document, trial, template, head)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            append_literal_row(sheet, row)
    payload = deterministic_workbook_bytes(workbook, creator="bctc-ai-f37-canonical-proposal")
    receipt = verify_f37_canonical_workbook_v1(payload, document, trial, template, head=head)
    return payload, receipt


def verify_f37_canonical_workbook_v1(payload: bytes, document: dict[str, Any], trial: dict[str, Any],
                                     template: list[tuple[int, str]], *, head: str) -> dict[str, Any]:
    """Read the workbook back and require every sheet cell to equal the projection exactly."""

    expected = _sheets_v1(document, trial, template, head)
    workbook = openpyxl.load_workbook(BytesIO(payload))
    try:
        _need(workbook.sheetnames == list(expected), "Workbook sheets differ")
        _need(not workbook_has_formula(workbook), "Workbook contains a formula")
        for name, rows in expected.items():
            actual = [list(row) for row in workbook[name].iter_rows(values_only=True)]
            actual = [row + [None] * (len(rows[0]) - len(row)) for row in actual]
            _need(actual == rows, f"Sheet {name} does not read back exactly")
    finally:
        workbook.close()
    values = sum(1 for (_rnid, _lane), slot in _tm_slots(trial, {r for r, _ in template}).items() if slot[0] is not None)
    return {"format": FORMAT, "workbook_sha256": hashlib.sha256(payload).hexdigest(), "workbook_size_bytes": len(payload),
            "mappings": len(trial["mappings"]), "tm_cells_written": values, "product_label": PRODUCT_LABEL}
