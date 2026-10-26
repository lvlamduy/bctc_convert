"""Exact, literal TM workbook rendering. Producer authentication lives upstream.

This is a per-document renderer, not an extraction or family acceptance gate.
It uses the base repository's canonical XLSX primitives, without notes/PRIMARY
imports that are absent at the migration base.
"""

from __future__ import annotations

import json
from hashlib import sha256
from io import BytesIO
from typing import Any

import openpyxl

from bctc_ai.export.canonical_xlsx import (
    append_literal_row,
    deterministic_workbook_bytes,
    workbook_has_formula,
)

TM_TEMPLATE_SHA256 = "64589f7ee2b025ac0bd784c13e0b0ba0aff9d46456a83fa7fe2d68a7de375115"
LABEL = "EXPERIMENTAL_READONLY_TM_EXPORT_NOT_FAMILY_ACCEPTANCE"
CODEC = "CANONICAL_SIGNED_INTEGER_V1"
SUPPORTED_STATES = {"RAW_SIGNED_INTEGER", "DASH_ZERO", "BLANK_SOURCE_CELL"}


class TmWorkbookError(ValueError):
    """A projection cannot be rendered without changing its meaning."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise TmWorkbookError(message)


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True,
                       separators=(",", ":"), allow_nan=False) + "\n").encode("utf-8")


def _json_rows(kind: str, identifier: str, value: Any) -> list[list[Any]]:
    raw = json_bytes(value)[:-1].decode("utf-8")
    # At most 24,000 UTF-16 code units even for non-BMP source characters.
    parts = [raw[n:n + 12000] for n in range(0, len(raw), 12000)] or [""]
    digest = sha256(raw.encode("utf-8")).hexdigest()
    return [[kind, identifier, n, len(parts), digest, part]
            for n, part in enumerate(parts, 1)]


def template_rows_v1(payload: bytes) -> list[tuple[int, str]]:
    """Read the RNID/name axis from the exact pinned template; never edit it."""
    require(sha256(payload).hexdigest() == TM_TEMPLATE_SHA256, "TM template SHA mismatch")
    workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        require(not workbook_has_formula(workbook), "template contains a formula")
        rows = []
        seen = set()
        for row in workbook.active.iter_rows(min_row=2, values_only=True):
            rnid, name = row[1:3] if len(row) >= 3 else (None, None)
            if rnid is None:
                require(name is None, "template name without RNID")
                continue
            require(type(rnid) is int and rnid > 0 and rnid not in seen,
                    "template RNID type or uniqueness drift")
            require(type(name) is str and bool(name), "template name missing")
            seen.add(rnid)
            rows.append((rnid, name))
        require(bool(rows), "template RNID axis is empty")
        return rows
    finally:
        workbook.close()


def exact_excel_cell_v1(cell: dict[str, Any]) -> tuple[Any, str]:
    """Keep dash/blank distinct and never round an integer to fit Excel."""
    require(type(cell) is dict and cell.get("state") in SUPPORTED_STATES,
            "unsupported transformed, derived, or unknown cell state")
    state, value, raw = (cell.get(key) for key in ("state", "coefficient", "source_text"))
    if state == "BLANK_SOURCE_CELL":
        require(value is None and raw is None, "blank has a fabricated value")
        return None, "TYPED_BLANK"
    require(type(value) is int and type(raw) is str and bool(raw), "exact source integer required")
    if state == "DASH_ZERO":
        require(value == 0 and raw.strip() in {"-", "–", "—", "_"}, "dash/source mismatch")
        return raw, "PRINTED_DASH_LITERAL"
    # Excel's decimal precision is 15 significant digits. Conservatively keep
    # every longer integer as exact decimal text, even if binary64 can hold it.
    if len(str(abs(value))) > 15:
        return str(value), "EXACT_INTEGER_TEXT_OVER_15_DIGITS"
    return value, "EXACT_INTEGER_NUMERIC"


def projected_sheet_rows_v1(
    *, document: dict[str, Any], records: list[dict[str, Any]],
    template: list[tuple[int, str]], provenance: dict[str, Any],
) -> dict[str, list[list[Any]]]:
    rnids = set()
    for rnid, name in template:
        require(type(rnid) is int and rnid > 0 and rnid not in rnids and type(name) is str,
                "invalid template row")
        rnids.add(rnid)
    require(bool(template), "empty template")
    periods, slots, ids = {}, {}, set()
    source_rows = [["kind", "id", "chunk", "chunks", "json_no_lf_sha256", "json_chunk"]]
    source_rows += _json_rows("DOCUMENT", "document", document)
    source_rows += _json_rows("AUTHENTICATED_INPUTS", "inputs", provenance)
    source_rows += _json_rows("TEMPLATE_ROWS", "template", template)
    for record in records:
        require(record["record_id"] not in ids, "duplicate projection record ID")
        ids.add(record["record_id"])
        require(record["document_id"] == document["document_id"], "foreign projection document")
        require(record["report_norm_id"] in rnids, "RNID is outside pinned template")
        require(record["coefficient_codec"] == CODEC, "unsupported coefficient codec")
        require(record["source_refs"] and record["period_proofs"] and record["unit_proofs"],
                "projection lacks source/period/unit proof")
        period = record["period"]
        key = json_bytes(period).decode("utf-8").strip()
        periods[key] = period
        slot = (record["report_norm_id"], key)
        require(slot not in slots, "duplicate document/RNID/full-period slot, even if values agree")
        slots[slot] = record
        source_rows += _json_rows("RECORD", record["record_id"], record)
    keys = sorted(periods)
    header = ["ReportNormId", "Name"]
    for key in keys:
        header += [f"{key}::{field}" for field in
                   ("Value", "State", "Unit", "RecordId", "NumberPolicy")]
    tm = [header]
    for rnid, name in template:
        row: list[Any] = [rnid, name]
        for key in keys:
            record = slots.get((rnid, key))
            if record is None:
                row += [None, "NO_ACCEPTED_MAPPING_FOR_SLOT", None, None, None]
            else:
                value, policy = exact_excel_cell_v1(record["cell"])
                row += [value, record["cell"]["state"], record["unit"], record["record_id"], policy]
        tm.append(row)
    return {
        "TM": tm,
        "Provenance": source_rows,
        "Status": [["Label", LABEL], ["DocumentStatus", document["status"]],
                   ["ProjectionRecordCount", len(records)],
                   ["FinancialAcceptance", "NOT_CLAIMED"],
                   ["NumericPolicy", "NO_SCALE_CONVERSION_NO_SIGN_FLIP_NO_IMPUTATION"]],
    }


def verify_projected_workbook_v1(payload: bytes, expected: dict[str, list[list[Any]]]) -> None:
    """Compare serialized values AND cell types, including trailing blank cells."""
    workbook = openpyxl.load_workbook(BytesIO(payload), read_only=False, data_only=False)
    try:
        require(workbook.sheetnames == list(expected), "workbook sheet axis drift")
        require(not workbook_has_formula(workbook), "workbook contains a formula")
        for name, rows in expected.items():
            sheet = workbook[name]
            require(sheet.max_row == len(rows), "workbook row axis drift")
            width = max(map(len, rows))
            require(sheet.max_column == width, "workbook column axis drift")
            for r, row in enumerate(rows, 1):
                for c in range(1, width + 1):
                    want = row[c - 1] if c <= len(row) else None
                    actual = sheet.cell(r, c)
                    # OpenPyXL reloads literal empty strings as None. This
                    # contract never emits an empty string; raw strings live
                    # inside canonical JSON chunks, so the distinction survives.
                    require(actual.value == want and type(actual.value) is type(want),
                            f"typed workbook value drift at {name}!R{r}C{c}")
                    if type(want) is str:
                        require(actual.data_type == "s", "text is not serialized as literal text")
                    elif type(want) is int:
                        require(actual.data_type == "n", "integer is not serialized as numeric")
    finally:
        workbook.close()


def render_projected_tm_workbook_v1(**kwargs: Any) -> tuple[bytes, dict[str, Any]]:
    rows = projected_sheet_rows_v1(**kwargs)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        for name, values in rows.items():
            require(len(values) <= 1048576 and max(map(len, values)) <= 16384,
                    "projection exceeds Excel row/column limits")
            sheet = workbook.create_sheet(name)
            for row in values:
                append_literal_row(sheet, row)
        require(not workbook_has_formula(workbook), "renderer introduced a formula")
        payload = deterministic_workbook_bytes(workbook, creator="BCTC_GEMINI_TM_BRIDGE_V1")
    finally:
        workbook.close()
    verify_projected_workbook_v1(payload, rows)
    return payload, {"label": LABEL, "record_count": len(kwargs["records"]),
                     "typed_readback": "PASS", "sheet_names": list(rows),
                     "sha256": sha256(payload).hexdigest(), "size_bytes": len(payload)}
